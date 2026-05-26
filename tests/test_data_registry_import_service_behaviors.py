from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app import create_app


class SaveMixin:
    def __init__(self):
        self.saved = []
        self.next_id = 1

    def save(self, obj):
        if getattr(obj, 'id', None) is None:
            obj.id = self.next_id
            self.next_id += 1
        self.saved.append(obj)
        return obj


class StubDataRegistryVersionRepository:
    def __init__(self, version=None):
        self.version = version
        self.saved = []

    def get_by_id(self, version_id):
        if self.version and self.version.id == version_id:
            return self.version
        return None

    def save(self, version):
        self.version = version
        self.saved.append(version)
        return version


class StubImportBatchRepository(SaveMixin):
    def __init__(self):
        super().__init__()
        self.batch = None

    def get_by_id(self, batch_id):
        if self.batch and self.batch.id == batch_id:
            return self.batch
        return None

    def list_by_registry_version(self, registry_version_id):
        return [
            batch for batch in self.saved if getattr(batch, 'registry_version_id', None) == registry_version_id
        ]

    def save(self, obj):
        saved = super().save(obj)
        self.batch = saved
        return saved


class StubImportRowRepository(SaveMixin):
    def __init__(self):
        super().__init__()
        self.rows_by_batch = {}

    def bulk_create(self, rows):
        for row in rows:
            self.save(row)
            self.rows_by_batch.setdefault(row.import_batch_id, []).append(row)
        return rows

    def list_by_batch(self, import_batch_id, status=None):
        rows = list(self.rows_by_batch.get(import_batch_id, []))
        if status is not None:
            rows = [row for row in rows if row.status == status]
        return rows

    def count_by_batch_and_status(self, import_batch_id):
        counts = {}
        for row in self.rows_by_batch.get(import_batch_id, []):
            counts[row.status] = counts.get(row.status, 0) + 1
        return counts


def make_draft_version(**overrides):
    defaults = {
        'id': 31,
        'registry_id': 10,
        'version_number': 1,
        'status': 'draft',
        'schema_json': {
            'fields': [
                {'key': 'record_code', 'type': 'string', 'required': True},
                {'key': 'label', 'type': 'string', 'required': True},
                {'key': 'effective_date', 'type': 'date', 'required': False},
            ]
        },
    }
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


def create_services(version=None):
    from app.modules.data_registry.services import (
        DataRegistryImportBatchService,
        DataRegistryImportValidationService,
    )

    batch_repository = StubImportBatchRepository()
    row_repository = StubImportRowRepository()
    version_repository = StubDataRegistryVersionRepository(version=version or make_draft_version())

    batch_service = DataRegistryImportBatchService(
        batch_repository=batch_repository,
        row_repository=row_repository,
        version_repository=version_repository,
    )
    validation_service = DataRegistryImportValidationService(
        batch_repository=batch_repository,
        row_repository=row_repository,
        version_repository=version_repository,
    )
    return batch_service, validation_service, batch_repository, row_repository


def create_batch_payload(rows=None):
    return {
        'batch_type': 'file_upload',
        'original_filename': 'program-2026.xlsx',
        'mime_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'reporting_year': 2026,
        'mapping_snapshot': {
            'fields': {
                'record_code': {'source': 'Kode Program'},
                'label': {'source': 'Nama Program'},
                'effective_date': {
                    'source': 'Tanggal Berlaku',
                    'accepted_input_formats': ['%d/%m/%Y', '%Y-%m-%d'],
                    'target_format': '%Y-%m-%d',
                },
            }
        },
        'source_headers': ['Kode Program', 'Nama Program', 'Tanggal Berlaku'],
        'source_snapshot': {'source_name': 'template-import-program'},
        'rows': rows
        or [
            {
                'Kode Program': 'PRG-001',
                'Nama Program': 'Program A',
                'Tanggal Berlaku': '21/05/2026',
            }
        ],
    }


def build_mapping_workbook_upload(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Sheet1'
    for row in rows:
        worksheet.append(row)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    payload = buffer.getvalue()
    return SimpleNamespace(
        filename='registry-import.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        read=lambda: payload,
    )


def test_get_batch_detail_returns_batch_and_version_context():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryImportBatchService

        version = make_draft_version(
            mapping_spec={'materialization_contract': 'wilayah_v1'},
            materialization_metadata={'batch_id': 101, 'record_count': 4},
            materialized_at='2026-05-22T11:30:00+00:00',
            freshness_status='fresh',
        )
        batch_service, _, batch_repository, row_repository = create_services(version=version)
        batch = batch_service.create_batch(31, create_batch_payload())['batch']
        batch.status = 'materialized'
        batch.materialization_summary = {'record_count': 4, 'source_row_count': 1}

        service = DataRegistryImportBatchService(
            batch_repository=batch_repository,
            row_repository=row_repository,
            version_repository=StubDataRegistryVersionRepository(version=version),
        )

        result = service.get_batch_detail(batch.id)

        assert result['batch'].id == batch.id
        assert result['version'].id == 31
        assert result['version'].freshness_status == 'fresh'
        assert result['version'].materialization_metadata == {'batch_id': 101, 'record_count': 4}


def test_create_batch_rejects_non_draft_target_version():
    app = create_app('testing')

    with app.app_context():
        batch_service, _validation_service, _batch_repository, _row_repository = create_services(
            version=make_draft_version(status='published')
        )

        with pytest.raises(ValueError, match='draft'):
            batch_service.create_batch(31, create_batch_payload())


def test_create_batch_persists_batch_and_staged_rows():
    app = create_app('testing')

    with app.app_context():
        batch_service, _validation_service, batch_repository, row_repository = create_services()
        actor = SimpleNamespace(id=7, uuid='actor-uuid')

        result = batch_service.create_batch(31, create_batch_payload(), actor=actor)

        assert result['batch'].status == 'mapped'
        assert result['batch'].total_rows == 1
        assert result['batch'].mapped_rows == 1
        assert batch_repository.batch.created_by == 7
        assert len(row_repository.list_by_batch(result['batch'].id)) == 1
        assert row_repository.list_by_batch(result['batch'].id)[0].status == 'mapped'


def test_validate_batch_marks_row_valid_and_normalizes_date():
    app = create_app('testing')

    with app.app_context():
        batch_service, validation_service, _batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(31, create_batch_payload())['batch']

        result = validation_service.validate_batch(batch.id)
        rows = row_repository.list_by_batch(batch.id)

        assert result['batch'].status == 'validated'
        assert result['batch'].valid_rows == 1
        assert rows[0].status == 'valid'
        assert rows[0].normalized_payload['effective_date'] == '2026-05-21'


def test_validate_batch_marks_required_field_errors_explicitly():
    app = create_app('testing')

    with app.app_context():
        batch_service, validation_service, _batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(
            31,
            create_batch_payload(
                rows=[
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': '',
                        'Tanggal Berlaku': '21/05/2026',
                    }
                ]
            ),
        )['batch']

        result = validation_service.validate_batch(batch.id)
        rows = row_repository.list_by_batch(batch.id)

        assert result['batch'].error_rows == 1
        assert rows[0].status == 'error'
        assert 'label' in rows[0].validation_errors


def test_validate_batch_rejects_invalid_date_format_explicitly():
    app = create_app('testing')

    with app.app_context():
        batch_service, validation_service, _batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(
            31,
            create_batch_payload(
                rows=[
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': 'Program A',
                        'Tanggal Berlaku': '31/02/2026',
                    }
                ]
            ),
        )['batch']

        result = validation_service.validate_batch(batch.id)
        rows = row_repository.list_by_batch(batch.id)

        assert result['batch'].error_rows == 1
        assert rows[0].status == 'error'
        assert 'effective_date' in rows[0].validation_errors


def test_validate_batch_rejects_invalid_select_option_explicitly():
    app = create_app('testing')

    with app.app_context():
        version = make_draft_version(
            schema_json={
                'fields': [
                    {'key': 'record_code', 'type': 'string', 'required': True},
                    {'key': 'label', 'type': 'string', 'required': True},
                    {
                        'key': 'status',
                        'type': 'select',
                        'required': True,
                        'options': ['aktif', 'nonaktif'],
                    },
                ]
            }
        )
        batch_service, validation_service, _batch_repository, row_repository = create_services(version=version)
        payload = create_batch_payload(
            rows=[
                {
                    'Kode Program': 'PRG-001',
                    'Nama Program': 'Program A',
                    'Status': 'draft',
                }
            ]
        )
        payload['mapping_snapshot']['fields']['status'] = {'source': 'Status'}
        payload['source_headers'] = ['Kode Program', 'Nama Program', 'Status']

        batch = batch_service.create_batch(31, payload)['batch']
        result = validation_service.validate_batch(batch.id)
        rows = row_repository.list_by_batch(batch.id)

        assert result['batch'].error_rows == 1
        assert rows[0].status == 'error'
        assert rows[0].validation_errors['status'] == 'invalid_option'


def test_validate_batch_marks_duplicate_record_code_within_batch():
    app = create_app('testing')

    with app.app_context():
        batch_service, validation_service, _batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(
            31,
            create_batch_payload(
                rows=[
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': 'Program A',
                        'Tanggal Berlaku': '21/05/2026',
                    },
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': 'Program B',
                        'Tanggal Berlaku': '2026-05-22',
                    },
                ]
            ),
        )['batch']

        result = validation_service.validate_batch(batch.id)
        rows = row_repository.list_by_batch(batch.id)

        assert result['batch'].duplicate_rows == 1
        assert [row.status for row in rows] == ['valid', 'duplicate']
        assert rows[1].duplicate_of_row_id == rows[0].id


def test_materialize_import_batch_rejects_batch_that_is_not_validated():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryMaterializationService

        batch_service, _validation_service, batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(31, create_batch_payload())['batch']
        service = DataRegistryMaterializationService(
            version_repository=StubDataRegistryVersionRepository(version=make_draft_version()),
            batch_repository=batch_repository,
            row_repository=row_repository,
        )

        with pytest.raises(ValueError, match='validated'):
            service.materialize_import_batch(batch.id)


def test_materialize_import_batch_allows_re_materialization_for_existing_materialized_batch(monkeypatch):
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryMaterializationService

        version = make_draft_version(
            mapping_spec={'materialization_contract': 'generic_v1'},
            source_snapshot={'source_name': 'diskominfo-jabar'},
        )
        batch_service, validation_service, batch_repository, row_repository = create_services(version=version)
        batch = batch_service.create_batch(
            31,
            create_batch_payload(
                rows=[
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': 'Program A',
                        'Tanggal Berlaku': '2026-01-15',
                    }
                ]
            ),
        )['batch']
        actor = SimpleNamespace(id=7, uuid='actor-uuid')
        validation_service.validate_batch(batch.id, actor=actor)

        service = DataRegistryMaterializationService(
            version_repository=StubDataRegistryVersionRepository(version=version),
            batch_repository=batch_repository,
            row_repository=row_repository,
        )
        calls = []

        def fake_materialize_wilayah_rows(registry_version_id, source_rows, actor=None):
            rows = list(source_rows)
            calls.append(
                {
                    'registry_version_id': registry_version_id,
                    'source_rows': rows,
                    'actor': actor,
                }
            )
            return {
                'registry': None,
                'version': version,
                'record_count': len(rows),
                'source_row_count': len(rows),
            }

        monkeypatch.setattr(service, 'materialize_wilayah_rows', fake_materialize_wilayah_rows)

        first_result = service.materialize_import_batch(batch.id, actor=actor)
        second_result = service.materialize_import_batch(batch.id, actor=actor)

        assert len(calls) == 2
        assert calls[0]['registry_version_id'] == batch.registry_version_id
        assert calls[1]['registry_version_id'] == batch.registry_version_id
        assert calls[0]['source_rows'] == calls[1]['source_rows']
        assert first_result['batch'].status == 'materialized'
        assert second_result['batch'].status == 'materialized'
        assert second_result['batch'].materialization_summary == {
            'record_count': 1,
            'source_row_count': 1,
            'materialization_contract': 'generic_v1',
        }
        assert second_result['version'].materialization_metadata['batch_id'] == batch.id
        assert second_result['version'].materialization_metadata['record_count'] == 1


def test_materialize_import_batch_forwards_canonical_wilayah_payload_and_persists_metadata(monkeypatch):
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryMaterializationService

        version = make_draft_version(
            registry_id=10,
            schema_json={
                'fields': [
                    {'key': 'import_row_id', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_provinsi_kode', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_kota_kode', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_kecamatan_kode', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_kelurahan_kode', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_provinsi_nama', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_kota_nama', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_kecamatan_nama', 'type': 'string', 'required': True},
                    {'key': 'kemendagri_kelurahan_nama', 'type': 'string', 'required': True},
                    {'key': 'bps_provinsi_kode', 'type': 'string', 'required': False},
                    {'key': 'bps_kota_kode', 'type': 'string', 'required': False},
                    {'key': 'bps_kecamatan_kode', 'type': 'string', 'required': False},
                    {'key': 'bps_kelurahan_kode', 'type': 'string', 'required': False},
                    {'key': 'bps_provinsi_nama', 'type': 'string', 'required': False},
                    {'key': 'bps_kota_nama', 'type': 'string', 'required': False},
                    {'key': 'bps_kecamatan_nama', 'type': 'string', 'required': False},
                    {'key': 'bps_kelurahan_nama', 'type': 'string', 'required': False},
                    {'key': 'latitude', 'type': 'string', 'required': False},
                    {'key': 'longitude', 'type': 'string', 'required': False},
                    {'key': 'kode_pos', 'type': 'string', 'required': False},
                    {'key': 'status_adm', 'type': 'string', 'required': False},
                ]
            },
            mapping_spec={
                'materialization_contract': 'wilayah_v1',
            },
            source_snapshot={'source_name': 'diskominfo-jabar'},
        )
        batch_service, validation_service, batch_repository, row_repository = create_services(version=version)
        batch = batch_service.create_batch(
            31,
            {
                'batch_type': 'file_upload',
                'original_filename': 'wilayah-jabar.xlsx',
                'mime_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                'reporting_year': 2026,
                'mapping_snapshot': {
                    'materialization_contract': 'wilayah_v1',
                    'fields': {
                        'import_row_id': {'source': 'ID'},
                        'kemendagri_provinsi_kode': {'source': 'Kode Prov Kemendagri'},
                        'kemendagri_kota_kode': {'source': 'Kode Kab/Kota Kemendagri'},
                        'kemendagri_kecamatan_kode': {'source': 'Kode Kecamatan Kemendagri'},
                        'kemendagri_kelurahan_kode': {'source': 'Kode Kelurahan Kemendagri'},
                        'kemendagri_provinsi_nama': {'source': 'Nama Prov Kemendagri'},
                        'kemendagri_kota_nama': {'source': 'Nama Kab/Kota Kemendagri'},
                        'kemendagri_kecamatan_nama': {'source': 'Nama Kecamatan Kemendagri'},
                        'kemendagri_kelurahan_nama': {'source': 'Nama Kelurahan Kemendagri'},
                        'bps_provinsi_kode': {'source': 'Kode Prov BPS'},
                        'bps_kota_kode': {'source': 'Kode Kab/Kota BPS'},
                        'bps_kecamatan_kode': {'source': 'Kode Kecamatan BPS'},
                        'bps_kelurahan_kode': {'source': 'Kode Kelurahan BPS'},
                        'bps_provinsi_nama': {'source': 'Nama Prov BPS'},
                        'bps_kota_nama': {'source': 'Nama Kab/Kota BPS'},
                        'bps_kecamatan_nama': {'source': 'Nama Kecamatan BPS'},
                        'bps_kelurahan_nama': {'source': 'Nama Kelurahan BPS'},
                        'latitude': {'source': 'Latitude'},
                        'longitude': {'source': 'Longitude'},
                        'kode_pos': {'source': 'Kode Pos'},
                        'status_adm': {'source': 'Status Administrasi'},
                    },
                },
                'source_headers': [],
                'source_snapshot': {'source_name': 'diskominfo-jabar'},
                'rows': [
                    {
                        'ID': '1',
                        'Kode Prov Kemendagri': '32',
                        'Kode Kab/Kota Kemendagri': '32.01',
                        'Kode Kecamatan Kemendagri': '32.01.01',
                        'Kode Kelurahan Kemendagri': '32.01.01.1001',
                        'Nama Prov Kemendagri': 'JAWA BARAT',
                        'Nama Kab/Kota Kemendagri': 'KAB. BOGOR',
                        'Nama Kecamatan Kemendagri': 'CIBINONG',
                        'Nama Kelurahan Kemendagri': 'HARAPANJAYA',
                        'Kode Prov BPS': '32.0',
                        'Kode Kab/Kota BPS': '3201.0',
                        'Kode Kecamatan BPS': '3201010.0',
                        'Kode Kelurahan BPS': '3201010001.0',
                        'Nama Prov BPS': 'JAWA BARAT',
                        'Nama Kab/Kota BPS': 'KABUPATEN BOGOR',
                        'Nama Kecamatan BPS': 'CIBINONG',
                        'Nama Kelurahan BPS': 'HARAPAN JAYA',
                        'Latitude': '-6.485088',
                        'Longitude': '106.854729',
                        'Kode Pos': '16913.0',
                        'Status Administrasi': '',
                    }
                ],
            },
        )['batch']
        actor = SimpleNamespace(id=7, uuid='actor-uuid')
        validation_service.validate_batch(batch.id, actor=actor)

        service = DataRegistryMaterializationService(
            version_repository=StubDataRegistryVersionRepository(version=version),
            batch_repository=batch_repository,
            row_repository=row_repository,
        )
        captured = {}

        def fake_materialize_wilayah_rows(registry_version_id, source_rows, actor=None):
            rows = list(source_rows)
            captured['registry_version_id'] = registry_version_id
            captured['source_rows'] = rows
            captured['actor'] = actor
            return {
                'registry': None,
                'version': version,
                'record_count': len(rows) * 4,
                'source_row_count': len(rows),
            }

        monkeypatch.setattr(service, 'materialize_wilayah_rows', fake_materialize_wilayah_rows)

        result = service.materialize_import_batch(batch.id, actor=actor)

        assert captured['registry_version_id'] == 31
        assert captured['actor'] is actor
        assert captured['source_rows'] == [
            {
                'id': '1',
                'kemendagri_provinsi_kode': '32',
                'kemendagri_kota_kode': '32.01',
                'kemendagri_kecamatan_kode': '32.01.01',
                'kemendagri_kelurahan_kode': '32.01.01.1001',
                'kemendagri_provinsi_nama': 'JAWA BARAT',
                'kemendagri_kota_nama': 'KAB. BOGOR',
                'kemendagri_kecamatan_nama': 'CIBINONG',
                'kemendagri_kelurahan_nama': 'HARAPANJAYA',
                'bps_provinsi_kode': '32.0',
                'bps_kota_kode': '3201.0',
                'bps_kecamatan_kode': '3201010.0',
                'bps_kelurahan_kode': '3201010001.0',
                'bps_provinsi_nama': 'JAWA BARAT',
                'bps_kota_nama': 'KABUPATEN BOGOR',
                'bps_kecamatan_nama': 'CIBINONG',
                'bps_kelurahan_nama': 'HARAPAN JAYA',
                'latitude': '-6.485088',
                'longitude': '106.854729',
                'kode_pos': '16913.0',
                'status_adm': None,
            }
        ]
        assert result['batch'].status == 'materialized'
        assert result['batch'].materialized_at is not None
        assert result['batch'].materialized_by == 7
        assert result['batch'].materialized_by_uuid == 'actor-uuid'
        assert result['batch'].materialization_summary == {
            'record_count': 4,
            'source_row_count': 1,
            'materialization_contract': 'wilayah_v1',
        }
        assert result['version'].materialized_at is not None
        assert result['version'].materialization_metadata == {
            'materialization_contract': 'wilayah_v1',
            'source_row_count': 1,
            'record_count': 4,
            'batch_id': batch.id,
            'batch_uuid': batch.uuid,
            'source_name': 'diskominfo-jabar',
        }
        assert result['materialization']['materialization_contract'] == 'wilayah_v1'


def test_export_import_batch_errors_builds_excel_with_source_headers_and_error_notes():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryImportBatchService

        batch_service, validation_service, batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(
            31,
            create_batch_payload(
                rows=[
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': 'Program Acuan',
                        'Tanggal Berlaku': '2026-05-21',
                    },
                    {
                        'Kode Program': 'PRG-001',
                        'Nama Program': 'Program Duplikat',
                        'Tanggal Berlaku': '2026-05-21',
                    },
                    {
                        'Kode Program': 'PRG-002',
                        'Nama Program': '',
                        'Tanggal Berlaku': '2026-05-22',
                    },
                ]
            ),
        )['batch']
        validation_service.validate_batch(batch.id)

        service = DataRegistryImportBatchService(
            batch_repository=batch_repository,
            row_repository=row_repository,
            version_repository=StubDataRegistryVersionRepository(version=make_draft_version()),
        )
        result = service.export_import_batch_errors(batch.id)

        workbook = load_workbook(filename=BytesIO(result['content']), read_only=True, data_only=True)
        data_sheet = workbook[workbook.sheetnames[0]]
        meta_sheet = workbook['_meta']
        error_sheet = workbook['_errors']
        data_rows = list(data_sheet.iter_rows(values_only=True))
        meta_rows = list(meta_sheet.iter_rows(values_only=True))
        error_rows = list(error_sheet.iter_rows(values_only=True))
        workbook.close()

        assert result['filename'] == 'data-registry-10-v1-errors.xlsx'
        assert result['error_rows_count'] == 2
        assert data_rows[0] == ('Kode Program', 'Nama Program', 'Tanggal Berlaku')
        assert str(data_rows[1][0] or '').find('duplicate_record_code') >= 0
        assert str(data_rows[2][1] or '').find('required') >= 0
        assert meta_rows[1][0] == 'batch_uuid'
        assert meta_rows[1][1] == batch.uuid
        assert error_rows[0] == ('row_number', 'status', 'field', 'source_header', 'error_code', 'error_message', 'current_value')
        assert any(row[2] == 'label' and row[3] == 'Nama Program' for row in error_rows[1:])


def test_export_import_batch_errors_rejects_when_batch_has_no_error_rows():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryImportBatchService

        batch_service, validation_service, batch_repository, row_repository = create_services()
        batch = batch_service.create_batch(31, create_batch_payload())['batch']
        validation_service.validate_batch(batch.id)

        service = DataRegistryImportBatchService(
            batch_repository=batch_repository,
            row_repository=row_repository,
            version_repository=StubDataRegistryVersionRepository(version=make_draft_version()),
        )

        with pytest.raises(ValueError, match='row error atau duplikat'):
            service.export_import_batch_errors(batch.id)


def test_parse_mapping_workbook_reads_headers_sample_rows_and_auto_mapping():
    app = create_app('testing')

    with app.app_context():
        batch_service, _validation_service, _batch_repository, _row_repository = create_services()
        upload = build_mapping_workbook_upload([
            ['Kode Program', 'Nama Program', 'Tanggal Berlaku'],
            ['PRG-001', 'Program A', '2026-05-21'],
            ['PRG-002', 'Program B', '2026-05-22'],
        ])

        result = batch_service.parse_mapping_workbook(31, upload)

        assert result['filename'] == 'registry-import.xlsx'
        assert result['headers'] == ['Kode Program', 'Nama Program', 'Tanggal Berlaku']
        assert result['sample_rows'][0]['Kode Program'] == 'PRG-001'
        assert result['sample_rows'][1]['Nama Program'] == 'Program B'
        assert result['auto_mapping']['record_code'] == 'Kode Program'
        assert result['auto_mapping']['label'] == 'Nama Program'
        assert result['auto_mapping']['effective_date'] == 'Tanggal Berlaku'
        assert result['fields'][0]['key'] == 'record_code'


def test_create_batch_from_workbook_builds_staged_rows_and_mapping_snapshot():
    app = create_app('testing')

    with app.app_context():
        batch_service, _validation_service, batch_repository, row_repository = create_services()
        actor = SimpleNamespace(id=9, uuid='actor-uuid')
        upload = build_mapping_workbook_upload([
            ['Kode Program', 'Nama Program', 'Tanggal Berlaku'],
            ['PRG-001', 'Program A', '2026-05-21'],
            ['PRG-002', 'Program B', '2026-05-22'],
        ])

        result = batch_service.create_batch_from_workbook(
            31,
            upload,
            {
                'record_code': 'Kode Program',
                'label': 'Nama Program',
                'effective_date': 'Tanggal Berlaku',
            },
            actor=actor,
        )

        batch = result['batch']
        rows = row_repository.list_by_batch(batch.id)

        assert batch.original_filename == 'registry-import.xlsx'
        assert batch.source_headers == ['Kode Program', 'Nama Program', 'Tanggal Berlaku']
        assert batch.mapping_snapshot['fields']['record_code']['source'] == 'Kode Program'
        assert batch.mapping_snapshot['fields']['label']['label'] == 'label'
        assert batch.total_rows == 2
        assert batch.created_by == 9
        assert batch_repository.batch.id == batch.id
        assert len(rows) == 2
        assert rows[0].mapped_payload['record_code'] == 'PRG-001'
        assert rows[1].mapped_payload['label'] == 'Program B'


def test_create_manual_entry_batch_builds_single_staged_row_and_mapping_snapshot():
    app = create_app('testing')

    with app.app_context():
        version = make_draft_version(schema_json={
            'fields': [
                {'key': 'record_key', 'type': 'string', 'required': True},
                {'key': 'record_code', 'type': 'string', 'required': True},
                {'key': 'label', 'type': 'string', 'required': True},
                {'key': 'parent_code', 'type': 'string', 'required': False},
                {'key': 'admin_level', 'type': 'string', 'required': False},
            ]
        })
        batch_service, _validation_service, batch_repository, row_repository = create_services(version=version)
        actor = SimpleNamespace(id=9, uuid='actor-uuid')

        result = batch_service.create_manual_entry_batch(
            31,
            {
                'record_key': 'program:prg-001',
                'record_code': 'PRG-001',
                'label': 'Program A',
                'parent_code': '',
                'admin_level': 'program',
            },
            actor=actor,
        )

        batch = result['batch']
        rows = row_repository.list_by_batch(batch.id)

        assert batch.batch_type == 'manual_entry'
        assert batch.original_filename == 'manual-entry-v31.xlsx'
        assert batch.source_headers == ['record_key', 'record_code', 'label', 'parent_code', 'admin_level']
        assert batch.mapping_snapshot['fields']['record_key']['source'] == 'record_key'
        assert batch.mapping_snapshot['fields']['record_code']['source'] == 'record_code'
        assert batch.source_snapshot['created_from'] == 'browser_manual_entry'
        assert batch.total_rows == 1
        assert batch.created_by == 9
        assert batch_repository.batch.id == batch.id
        assert len(rows) == 1
        assert rows[0].raw_payload['record_code'] == 'PRG-001'
        assert rows[0].mapped_payload['label'] == 'Program A'
