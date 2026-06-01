from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from app import create_app
from app.modules.data_registry.models import DataRegistry, DataRegistryVersion


class SaveMixin:
    def __init__(self):
        self.saved = []
        self.next_id = 1

    def save(self, obj):
        if getattr(obj, 'id', None) is None:
            obj.id = self.next_id
            self.next_id += 1
        self.saved.append(obj)
        return obj


class StubDataRegistryRepository(SaveMixin):
    def __init__(self, registry=None):
        super().__init__()
        self.registry = registry

    def get_by_id(self, registry_id):
        if self.registry and self.registry.id == registry_id:
            return self.registry
        return None

    def get_by_slug(self, slug):
        if self.registry and self.registry.registry_slug == slug:
            return self.registry
        return None

    def get_active_by_slug(self, slug):
        if self.registry and self.registry.registry_slug == slug and self.registry.status in {'draft', 'published'}:
            return self.registry
        return None

    def find_by(self, **filters):
        if self.registry and filters.get('registry_code') and self.registry.registry_code == filters['registry_code']:
            return [self.registry]
        return []


class StubDataRegistryVersionRepository(SaveMixin):
    def __init__(self, version=None):
        super().__init__()
        self.version = version
        self.archived_args = None

    def get_by_id(self, version_id):
        if self.version and self.version.id == version_id:
            return self.version
        return None

    def get_next_version_number(self, registry_id):
        return 2

    def archive_published_others(self, registry_id, except_version_id=None):
        self.archived_args = {
            'registry_id': registry_id,
            'except_version_id': except_version_id,
        }
        return []

    def get_published_version(self, registry_id):
        if self.version and self.version.registry_id == registry_id and self.version.status == 'published':
            return self.version
        return None


class StubDraftRegistryVersionService:
    def __init__(self):
        self.called_with = None

    def create_draft_version(self, **kwargs):
        self.called_with = kwargs
        return DataRegistryVersion(
            id=11,
            registry_id=kwargs['registry_id'],
            version_number=1,
            schema_json=kwargs['schema_json'],
            mapping_spec=kwargs.get('mapping_spec') or {},
            source_snapshot=kwargs.get('source_snapshot') or {},
            status='draft',
            freshness_status='unknown',
        )


class StubQueryVersionRepository:
    def __init__(self, published_version=None, versions=None):
        self.published_version = published_version
        self.versions = versions or ([] if published_version is None else [published_version])
        self.published_calls = []
        self.list_calls = []

    def get_published_version(self, registry_id):
        self.published_calls.append(registry_id)
        if self.published_version and self.published_version.registry_id == registry_id:
            return self.published_version
        return None

    def list_versions(self, registry_id, include_deleted=False):
        self.list_calls.append({'registry_id': registry_id, 'include_deleted': include_deleted})
        return [version for version in self.versions if version.registry_id == registry_id]


class StubQueryRecordRepository:
    def __init__(self, records):
        self.records = records
        self.list_options_calls = []
        self.list_children_calls = []
        self.list_by_level_calls = []
        self.list_feature_collection_calls = []

    def get_by_key(self, registry_version_id, record_key):
        for record in self.records:
            if record.registry_version_id == registry_version_id and record.record_key == record_key:
                return record
        return None

    def get_by_code(self, registry_version_id, record_code):
        for record in self.records:
            if record.registry_version_id == registry_version_id and record.record_code == record_code:
                return record
        return None

    def list_children(self, registry_version_id, parent_record_id):
        self.list_children_calls.append(
            {'registry_version_id': registry_version_id, 'parent_record_id': parent_record_id}
        )
        items = [
            record
            for record in self.records
            if record.registry_version_id == registry_version_id and record.parent_record_id == parent_record_id
        ]
        return sorted(items, key=lambda item: ((item.sort_order or 0), item.label))

    def list_by_level(self, registry_version_id, admin_level, parent_record_id=None):
        self.list_by_level_calls.append(
            {
                'registry_version_id': registry_version_id,
                'admin_level': admin_level,
                'parent_record_id': parent_record_id,
            }
        )
        items = [
            record
            for record in self.records
            if record.registry_version_id == registry_version_id and record.admin_level == admin_level
        ]
        if parent_record_id is not None:
            items = [record for record in items if record.parent_record_id == parent_record_id]
        return sorted(items, key=lambda item: ((item.sort_order or 0), item.label))

    def list_options(self, registry_version_id, admin_level=None, parent_record_id=None, q=None, limit=100):
        self.list_options_calls.append(
            {
                'registry_version_id': registry_version_id,
                'admin_level': admin_level,
                'parent_record_id': parent_record_id,
                'q': q,
                'limit': limit,
            }
        )
        items = [record for record in self.records if record.registry_version_id == registry_version_id]
        if admin_level is not None:
            items = [record for record in items if record.admin_level == admin_level]
        if parent_record_id is not None:
            items = [record for record in items if record.parent_record_id == parent_record_id]
        if q:
            q_lower = q.lower()
            items = [record for record in items if q_lower in record.label.lower()]
        items = sorted(items, key=lambda item: ((item.sort_order or 0), item.label))
        return items[:limit]

    def list_feature_collection_records(self, registry_version_id, admin_level=None):
        self.list_feature_collection_calls.append(
            {'registry_version_id': registry_version_id, 'admin_level': admin_level}
        )
        items = [record for record in self.records if record.registry_version_id == registry_version_id]
        if admin_level is not None:
            items = [record for record in items if record.admin_level == admin_level]
        return sorted(items, key=lambda item: ((item.sort_order or 0), item.label))


class StubMaterializationRecordRepository(SaveMixin):
    def __init__(self):
        super().__init__()
        self.deleted_registry_version_id = None

    def delete_by_registry_version(self, registry_version_id):
        self.deleted_registry_version_id = registry_version_id
        self.saved = [item for item in self.saved if item.registry_version_id != registry_version_id]
        return True


def make_record(**overrides):
    defaults = {
        'id': 1,
        'registry_id': 10,
        'registry_version_id': 30,
        'parent_record_id': None,
        'record_key': 'province:32',
        'record_code': '32',
        'external_code': None,
        'label': 'Jawa Barat',
        'display_label': None,
        'normalized_label': 'jawa barat',
        'admin_level': 'province',
        'admin_level_code': 'PROV',
        'city_regency_kind': None,
        'province_code': '32',
        'city_regency_code': None,
        'district_code': None,
        'village_code': None,
        'village_adm_status': None,
        'sort_order': 1,
        'is_active': True,
        'valid_from': None,
        'valid_to': None,
        'centroid_lat': None,
        'centroid_lng': None,
        'bbox_min_lat': None,
        'bbox_min_lng': None,
        'bbox_max_lat': None,
        'bbox_max_lng': None,
        'geometry_json': None,
        'source_row_number': None,
        'source_row_hash': None,
        'source_snapshot': {},
        'payload': {},
    }
    defaults.update(overrides)
    return SimpleNamespace(**defaults)


def test_create_registry_returns_registry_and_initial_draft_version():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryService

        registry_repository = StubDataRegistryRepository()
        version_service = StubDraftRegistryVersionService()
        service = DataRegistryService(
            registry_repository=registry_repository,
            version_service=version_service,
        )
        actor = SimpleNamespace(id=7, uuid='actor-uuid')

        result = service.create_registry(
            {
                'registry_slug': 'wilayah.administratif',
                'registry_code': 'WILAYAH-ADM',
                'name': 'Wilayah Administratif',
                'description': 'Registry wilayah administratif Jawa Barat.',
                'schema_json': {'fields': [{'key': 'kode_wilayah'}]},
                'mapping_spec': {'source': 'csv'},
                'source_snapshot': {'source_name': 'diskominfo-jabar'},
            },
            actor=actor,
        )

        assert result['registry'].id == 1
        assert result['registry'].registry_slug == 'wilayah.administratif'
        assert result['registry'].registry_code == 'WILAYAH-ADM'
        assert result['registry'].name == 'Wilayah Administratif'
        assert result['registry'].created_by == 7
        assert result['draft_version'].registry_id == 1
        assert result['draft_version'].schema_json == {'fields': [{'key': 'kode_wilayah'}]}
        assert result['draft_version'].mapping_spec == {'source': 'csv'}
        assert version_service.called_with['actor'] is actor


def test_create_registry_rejects_duplicate_registry_code():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryService

        existing_registry = DataRegistry(
            id=99,
            registry_slug='existing.registry',
            registry_code='MASTER-PROGRAM',
            name='Existing Registry',
            status='draft',
        )
        registry_repository = StubDataRegistryRepository(registry=existing_registry)
        version_service = StubDraftRegistryVersionService()
        service = DataRegistryService(
            registry_repository=registry_repository,
            version_service=version_service,
        )

        try:
            service.create_registry(
                {
                    'registry_slug': 'master.program',
                    'registry_code': 'MASTER-PROGRAM',
                    'name': 'Master Program',
                    'schema_json': {'fields': [{'key': 'record_code'}]},
                }
            )
            assert False, 'Expected ValueError for duplicate registry code'
        except ValueError as error:
            assert str(error) == 'Registry code sudah digunakan.'

        assert version_service.called_with is None
        assert registry_repository.saved == []


def test_create_registry_accepts_master_data_hierarchical_contract():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryService

        registry_repository = StubDataRegistryRepository()
        version_service = StubDraftRegistryVersionService()
        service = DataRegistryService(
            registry_repository=registry_repository,
            version_service=version_service,
        )

        result = service.create_registry(
            {
                'registry_slug': 'master.program',
                'registry_code': 'MASTER-PROGRAM',
                'name': 'Master Program',
                'description': 'Registry master data bertingkat.',
                'registry_type': 'master_data',
                'category_key': 'master_data',
                'source_mode': 'import_file',
                'data_shape': 'hierarchical',
                'schema_json': {
                    'fields': [
                        {'key': 'record_key', 'type': 'text', 'required': True},
                        {'key': 'record_code', 'type': 'text', 'required': True},
                        {'key': 'label', 'type': 'text', 'required': True},
                    ]
                },
                'mapping_spec': {'materialization_contract': 'generic_v1'},
                'source_snapshot': {'source_name': 'browser_manual_setup'},
            }
        )

        assert result['registry'].registry_type == 'master_data'
        assert result['registry'].data_shape == 'hierarchical'
        assert result['draft_version'].schema_json['fields'][0]['key'] == 'record_key'
        assert version_service.called_with['mapping_spec']['materialization_contract'] == 'generic_v1'


def test_update_draft_schema_fields_rewrites_schema_json_fields_only():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryVersionService

        version = DataRegistryVersion(
            id=31,
            registry_id=10,
            version_number=2,
            status='draft',
            schema_json={
                'fields': [
                    {'key': 'record_code', 'type': 'text', 'required': True},
                ],
                'layout': {'mode': 'simple'},
            },
            mapping_spec={'materialization_contract': 'generic_v1'},
            source_snapshot={'source_name': 'browser_manual_setup'},
            freshness_status='unknown',
        )
        version_repository = StubDataRegistryVersionRepository(version=version)
        registry_repository = StubDataRegistryRepository(
            registry=DataRegistry(id=10, registry_slug='master.program', name='Master Program', status='draft')
        )
        service = DataRegistryVersionService(
            version_repository=version_repository,
            registry_repository=registry_repository,
        )

        saved_version = service.update_draft_schema_fields(
            31,
            [
                {'key': 'record_key', 'label': 'Record Key', 'type': 'text', 'required': True},
                {
                    'key': 'record_code',
                    'label': 'Record Code',
                    'type': 'select',
                    'required': True,
                    'options': ['aktif', 'nonaktif'],
                },
                {'key': 'label', 'label': 'Label', 'type': 'text', 'required': True},
                {'key': 'effective_date', 'label': 'Tanggal Berlaku', 'type': 'date', 'required': False},
            ],
        )

        assert saved_version.schema_json['layout'] == {'mode': 'simple'}
        assert [field['key'] for field in saved_version.schema_json['fields']] == [
            'record_key',
            'record_code',
            'label',
            'effective_date',
        ]
        assert saved_version.schema_json['fields'][1]['options'] == ['aktif', 'nonaktif']
        assert saved_version.schema_json['fields'][3]['type'] == 'date'
        assert version_repository.saved[-1] is saved_version


def test_build_registry_template_workbook_uses_schema_fields_as_headers():
    app = create_app('testing')

    with app.app_context():
        from openpyxl import load_workbook
        from app.modules.data_registry.services import DataRegistryImportBatchService

        registry = DataRegistry(
            id=10,
            uuid='registry-uuid',
            registry_slug='master.program',
            registry_code='MASTER-PROGRAM',
            name='Master Program',
            status='draft',
        )
        version = DataRegistryVersion(
            id=31,
            uuid='version-uuid',
            registry_id=10,
            version_number=2,
            status='draft',
            schema_json={
                'fields': [
                    {'key': 'record_key', 'label': 'Record Key', 'type': 'text', 'required': True},
                    {
                        'key': 'record_code',
                        'label': 'Record Code',
                        'type': 'select',
                        'required': True,
                        'options': ['aktif', 'nonaktif'],
                    },
                    {'key': 'effective_date', 'label': 'Tanggal Berlaku', 'type': 'date', 'required': False},
                ]
            },
            freshness_status='unknown',
        )
        service = DataRegistryImportBatchService()

        workbook_bytes = service.build_template_workbook(version, registry=registry)
        workbook = load_workbook(filename=BytesIO(workbook_bytes))

        assert workbook.sheetnames == ['data', '_meta', '_dictionary']
        assert [cell.value for cell in workbook['data'][1]] == ['Record Key', 'Record Code', 'Tanggal Berlaku']
        assert workbook['_meta']['A11'].value == 'field_key'
        assert workbook['_meta']['B12'].value == 'Record Key'
        assert workbook['_meta']['E13'].value == 'aktif, nonaktif'
        assert workbook['_dictionary']['A2'].value == 'record_code'
        assert workbook['_dictionary']['C2'].value == 'aktif'
        assert service.build_template_filename(version, registry=registry) == 'MASTER-PROGRAM-v2-template.xlsx'


def test_create_draft_version_can_clone_from_source_version_when_schema_not_provided():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryVersionService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='draft',
        )
        source_version = DataRegistryVersion(
            id=20,
            registry_id=10,
            version_number=1,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            mapping_spec={'source': 'csv'},
            source_snapshot={'source_name': 'diskominfo-jabar'},
            status='published',
            freshness_status='fresh',
        )
        version_repository = StubDataRegistryVersionRepository(version=source_version)
        registry_repository = StubDataRegistryRepository(registry=registry)
        service = DataRegistryVersionService(
            version_repository=version_repository,
            registry_repository=registry_repository,
        )

        draft_version = service.create_draft_version(
            registry_id=10,
            schema_json=None,
            source_version_id=20,
        )

        assert draft_version.registry_id == 10
        assert draft_version.version_number == 2
        assert draft_version.schema_json == {'fields': [{'key': 'kode_wilayah'}]}
        assert draft_version.mapping_spec == {'source': 'csv'}
        assert draft_version.source_snapshot == {'source_name': 'diskominfo-jabar'}
        assert draft_version.status == 'draft'
        assert draft_version.freshness_status == 'unknown'


def test_publish_version_marks_version_published_and_parent_registry_published():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryVersionService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='draft',
        )
        version = DataRegistryVersion(
            id=20,
            registry_id=10,
            version_number=1,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            mapping_spec={'source': 'csv'},
            source_snapshot={'source_name': 'diskominfo-jabar'},
            status='draft',
            freshness_status='unknown',
        )
        version_repository = StubDataRegistryVersionRepository(version=version)
        registry_repository = StubDataRegistryRepository(registry=registry)
        service = DataRegistryVersionService(
            version_repository=version_repository,
            registry_repository=registry_repository,
        )

        published_version = service.publish_version(20)

        assert version_repository.archived_args == {
            'registry_id': 10,
            'except_version_id': 20,
        }
        assert published_version.status == 'published'
        assert published_version.published_at is not None
        assert published_version.freshness_status == 'fresh'
        assert registry.status == 'published'
        assert registry_repository.saved[-1] is registry


def test_get_option_list_reads_only_published_version_and_supports_parent_code_filter():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='published',
        )
        published_version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=3,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='published',
            freshness_status='fresh',
        )
        draft_version = DataRegistryVersion(
            id=31,
            registry_id=10,
            version_number=4,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='draft',
            freshness_status='unknown',
        )
        record_repository = StubQueryRecordRepository(
            [
                make_record(id=1, registry_version_id=30, record_key='province:32', record_code='32', label='Jawa Barat', admin_level='province'),
                make_record(id=2, registry_version_id=30, parent_record_id=1, record_key='city:3204', record_code='3204', label='Kabupaten Bandung', admin_level='city_regency', city_regency_code='3204'),
                make_record(id=3, registry_version_id=31, parent_record_id=999, record_key='city:9999', record_code='9999', label='Draft City', admin_level='city_regency'),
            ]
        )
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(
                published_version=published_version,
                versions=[published_version, draft_version],
            ),
            record_repository=record_repository,
        )

        result = service.get_option_list(
            'wilayah.administratif',
            admin_level='city_regency',
            parent_code='32',
        )

        assert result['version'].id == 30
        assert result['items'] == [
            {
                'label': 'Kabupaten Bandung',
                'value': '3204',
                'meta': {
                    'record_key': 'city:3204',
                    'record_code': '3204',
                    'admin_level': 'city_regency',
                    'parent_record_id': 1,
                },
            }
        ]
        assert record_repository.list_options_calls[-1] == {
            'registry_version_id': 30,
            'admin_level': 'city_regency',
            'parent_record_id': 1,
            'q': None,
            'limit': 100,
        }



def test_get_option_list_raises_when_published_version_missing():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='draft',
        )
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(published_version=None, versions=[]),
            record_repository=StubQueryRecordRepository([]),
        )

        try:
            service.get_option_list('wilayah.administratif')
            assert False, 'Expected ValueError when published version is missing.'
        except ValueError as error:
            assert str(error) == 'Published version registry tidak ditemukan.'



def test_get_lookup_rejects_when_key_and_code_missing():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='published',
        )
        published_version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=3,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='published',
            freshness_status='fresh',
        )
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(published_version=published_version),
            record_repository=StubQueryRecordRepository([]),
        )

        try:
            service.get_lookup('wilayah.administratif')
            assert False, 'Expected ValueError when record key/code is missing.'
        except ValueError as error:
            assert str(error) == 'record_key atau record_code wajib diisi.'



def test_get_lookup_can_find_record_by_code():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='published',
        )
        published_version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=3,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='published',
            freshness_status='fresh',
        )
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(published_version=published_version),
            record_repository=StubQueryRecordRepository(
                [
                    make_record(
                        id=2,
                        registry_version_id=30,
                        parent_record_id=1,
                        record_key='city:3204',
                        record_code='3204',
                        label='Kabupaten Bandung',
                        admin_level='city_regency',
                        city_regency_code='3204',
                    )
                ]
            ),
        )

        result = service.get_lookup('wilayah.administratif', record_code='3204')

        assert result['record'] == {
            'id': 2,
            'record_key': 'city:3204',
            'record_code': '3204',
            'label': 'Kabupaten Bandung',
            'display_label': 'Kabupaten Bandung',
            'admin_level': 'city_regency',
            'parent_record_id': 1,
            'centroid': None,
            'payload': {},
        }



def test_get_children_returns_only_active_children_and_respects_admin_level():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='published',
        )
        published_version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=3,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='published',
            freshness_status='fresh',
        )
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(published_version=published_version),
            record_repository=StubQueryRecordRepository(
                [
                    make_record(id=1, registry_version_id=30, record_key='province:32', record_code='32', label='Jawa Barat', admin_level='province'),
                    make_record(id=2, registry_version_id=30, parent_record_id=1, record_key='city:3204', record_code='3204', label='Kabupaten Bandung', admin_level='city_regency', sort_order=2),
                    make_record(id=3, registry_version_id=30, parent_record_id=1, record_key='city:3273', record_code='3273', label='Kota Bandung', admin_level='city_regency', sort_order=1),
                    make_record(id=4, registry_version_id=30, parent_record_id=1, record_key='district:x', record_code='x', label='Inactive District', admin_level='district', is_active=False, sort_order=3),
                ]
            ),
        )

        result = service.get_children('wilayah.administratif', parent_code='32', admin_level='city_regency')

        assert [item['record_code'] for item in result['items']] == ['3273', '3204']
        assert [item['label'] for item in result['items']] == ['Kota Bandung', 'Kabupaten Bandung']



def test_get_tree_returns_nested_tree_until_max_depth():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='published',
        )
        published_version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=3,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='published',
            freshness_status='fresh',
        )
        records = [
            make_record(id=1, registry_version_id=30, record_key='province:32', record_code='32', label='Jawa Barat', admin_level='province'),
            make_record(id=2, registry_version_id=30, parent_record_id=1, record_key='city:3204', record_code='3204', label='Kabupaten Bandung', admin_level='city_regency'),
            make_record(id=3, registry_version_id=30, parent_record_id=2, record_key='district:3204010', record_code='3204010', label='Cicalengka', admin_level='district'),
            make_record(id=4, registry_version_id=30, parent_record_id=3, record_key='village:3204010001', record_code='3204010001', label='Babakan Peuteuy', admin_level='village'),
        ]
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(published_version=published_version),
            record_repository=StubQueryRecordRepository(records),
        )

        result = service.get_tree('wilayah.administratif', root_level='province', max_depth=3)

        assert result['items'][0]['record_code'] == '32'
        assert result['items'][0]['children'][0]['record_code'] == '3204'
        assert result['items'][0]['children'][0]['children'][0]['record_code'] == '3204010'
        assert result['items'][0]['children'][0]['children'][0]['children'] == []



def test_get_feature_collection_excludes_records_without_geometry_or_centroid():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryQueryService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='published',
        )
        published_version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=3,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='published',
            freshness_status='fresh',
        )
        records = [
            make_record(
                id=1,
                registry_version_id=30,
                record_key='province:32',
                record_code='32',
                label='Jawa Barat',
                admin_level='province',
                geometry_json={'type': 'Polygon', 'coordinates': []},
            ),
            make_record(
                id=2,
                registry_version_id=30,
                parent_record_id=1,
                record_key='city:3204',
                record_code='3204',
                label='Kabupaten Bandung',
                admin_level='city_regency',
                centroid_lat='-6.914744',
                centroid_lng='107.609810',
            ),
            make_record(
                id=3,
                registry_version_id=30,
                parent_record_id=1,
                record_key='city:missing',
                record_code='missing',
                label='No Geometry',
                admin_level='city_regency',
            ),
        ]
        service = DataRegistryQueryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=StubQueryVersionRepository(published_version=published_version),
            record_repository=StubQueryRecordRepository(records),
        )

        result = service.get_feature_collection('wilayah.administratif')

        assert result['type'] == 'FeatureCollection'
        assert len(result['features']) == 2
        assert result['features'][0]['properties']['record_code'] == '32'
        assert result['features'][1]['geometry'] == {
            'type': 'Point',
            'coordinates': [107.60981, -6.914744],
        }



def test_materialize_wilayah_rows_expands_flat_rows_into_hierarchy_and_refreshes_version():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryMaterializationService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='draft',
        )
        version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=1,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='draft',
            freshness_status='unknown',
        )
        record_repository = StubMaterializationRecordRepository()
        version_repository = StubDataRegistryVersionRepository(version=version)
        service = DataRegistryMaterializationService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=version_repository,
            record_repository=record_repository,
        )

        result = service.materialize_wilayah_rows(
            30,
            [
                {
                    'id': '1',
                    'kemendagri_provinsi_kode': '32',
                    'kemendagri_kota_kode': '32.01',
                    'kemendagri_kecamatan_kode': '32.01.01',
                    'kemendagri_kelurahan_kode': '32.01.01.1001',
                    'kemendagri_provinsi_nama': 'JAWA BARAT',
                    'kemendagri_kota_nama': 'KAB. BOGOR',
                    'kemendagri_kecamatan_nama': 'CIBINONG',
                    'kemendagri_kelurahan_nama': 'HARAPANJAYA',
                    'bps_provinsi_kode': '32.0',
                    'bps_kota_kode': '3201.0',
                    'bps_kecamatan_kode': '3201010.0',
                    'bps_kelurahan_kode': '3201010001.0',
                    'bps_provinsi_nama': 'JAWA BARAT',
                    'bps_kota_nama': 'KABUPATEN BOGOR',
                    'bps_kecamatan_nama': 'CIBINONG',
                    'bps_kelurahan_nama': 'HARAPAN JAYA',
                    'latitude': '-6.485088',
                    'longitude': '106.854729',
                    'kode_pos': '16913.0',
                    'status_adm': '',
                },
                {
                    'id': '2',
                    'kemendagri_provinsi_kode': '32',
                    'kemendagri_kota_kode': '32.01',
                    'kemendagri_kecamatan_kode': '32.01.01',
                    'kemendagri_kelurahan_kode': '32.01.01.1002',
                    'kemendagri_provinsi_nama': 'JAWA BARAT',
                    'kemendagri_kota_nama': 'KAB. BOGOR',
                    'kemendagri_kecamatan_nama': 'CIBINONG',
                    'kemendagri_kelurahan_nama': 'CIRIUNG',
                    'bps_provinsi_kode': '32.0',
                    'bps_kota_kode': '3201.0',
                    'bps_kecamatan_kode': '3201010.0',
                    'bps_kelurahan_kode': '3201010002.0',
                    'bps_provinsi_nama': 'JAWA BARAT',
                    'bps_kota_nama': 'KABUPATEN BOGOR',
                    'bps_kecamatan_nama': 'CIBINONG',
                    'bps_kelurahan_nama': 'CIRIUNG',
                    'latitude': '-6.490000',
                    'longitude': '106.860000',
                    'kode_pos': '16914.0',
                    'status_adm': '',
                },
            ],
        )

        assert record_repository.deleted_registry_version_id == 30
        assert result['record_count'] == 5
        assert len(record_repository.saved) == 5

        province, city, district, village_one, village_two = record_repository.saved
        assert province.record_key == 'province:32'
        assert province.record_code == '32'
        assert province.parent_record_id is None
        assert province.external_code == '32'
        assert province.payload['codes']['kemendagri']['province'] == '32'

        assert city.record_key == 'city_regency:32.01'
        assert city.parent_record_id == province.id
        assert city.city_regency_kind == 'kabupaten'
        assert city.external_code == '3201'

        assert district.record_key == 'district:32.01.01'
        assert district.parent_record_id == city.id
        assert district.external_code == '3201010'

        assert village_one.record_key == 'village:32.01.01.1001'
        assert village_one.parent_record_id == district.id
        assert village_one.external_code == '3201010001'
        assert village_one.village_adm_status == 'kelurahan'
        assert village_one.payload['postal_code'] == '16913'
        assert village_one.payload['source_labels']['bps']['village'] == 'HARAPAN JAYA'
        assert village_one.source_row_number == 1
        assert village_one.source_row_hash
        assert village_one.centroid_lat == Decimal('-6.485088')
        assert village_one.centroid_lng == Decimal('106.854729')

        assert village_two.record_code == '32.01.01.1002'
        assert village_two.parent_record_id == district.id
        assert village_two.payload['postal_code'] == '16914'

        assert version.freshness_status == 'fresh'
        assert version.materialized_watermark
        assert version.freshness_signature
        assert version_repository.saved[-1] is version


def test_materialize_wilayah_rows_nulls_duplicate_bps_external_codes_for_conflicting_records():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryMaterializationService

        registry = DataRegistry(
            id=10,
            registry_slug='wilayah.administratif',
            name='Wilayah Administratif',
            status='draft',
        )
        version = DataRegistryVersion(
            id=30,
            registry_id=10,
            version_number=1,
            schema_json={'fields': [{'key': 'kode_wilayah'}]},
            status='draft',
            freshness_status='unknown',
        )
        record_repository = StubMaterializationRecordRepository()
        version_repository = StubDataRegistryVersionRepository(version=version)
        service = DataRegistryMaterializationService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_repository=version_repository,
            record_repository=record_repository,
        )

        service.materialize_wilayah_rows(
            30,
            [
                {
                    'id': '1700',
                    'kemendagri_provinsi_kode': '32',
                    'kemendagri_kota_kode': '32.05',
                    'kemendagri_kecamatan_kode': '32.05.20',
                    'kemendagri_kelurahan_kode': '32.05.20.2003',
                    'kemendagri_provinsi_nama': 'JAWA BARAT',
                    'kemendagri_kota_nama': 'KAB. GARUT',
                    'kemendagri_kecamatan_nama': 'CISURUPAN',
                    'kemendagri_kelurahan_nama': 'SUKAWANGI',
                    'bps_provinsi_kode': '32.0',
                    'bps_kota_kode': '3205.0',
                    'bps_kecamatan_kode': '3205160.0',
                    'bps_kelurahan_kode': '3205160002.0',
                    'bps_provinsi_nama': 'JAWA BARAT',
                    'bps_kota_nama': 'KABUPATEN GARUT',
                    'bps_kecamatan_nama': 'CISURUPAN',
                    'bps_kelurahan_nama': 'SUKAWANGI',
                    'latitude': '-7.34350',
                    'longitude': '107.78944',
                    'kode_pos': '44163.0',
                    'status_adm': '',
                },
                {
                    'id': '1701',
                    'kemendagri_provinsi_kode': '32',
                    'kemendagri_kota_kode': '32.05',
                    'kemendagri_kecamatan_kode': '32.05.20',
                    'kemendagri_kelurahan_kode': '32.05.20.2004',
                    'kemendagri_provinsi_nama': 'JAWA BARAT',
                    'kemendagri_kota_nama': 'KAB. GARUT',
                    'kemendagri_kecamatan_nama': 'CISURUPAN',
                    'kemendagri_kelurahan_nama': 'SUKATANI',
                    'bps_provinsi_kode': '32.0',
                    'bps_kota_kode': '3205.0',
                    'bps_kecamatan_kode': '3205160.0',
                    'bps_kelurahan_kode': '3205160002.0',
                    'bps_provinsi_nama': 'JAWA BARAT',
                    'bps_kota_nama': 'KABUPATEN GARUT',
                    'bps_kecamatan_nama': 'CISURUPAN',
                    'bps_kelurahan_nama': 'SUKATANI',
                    'latitude': '-7.34351',
                    'longitude': '107.78945',
                    'kode_pos': '44163.0',
                    'status_adm': '',
                },
            ],
        )

        villages = [item for item in record_repository.saved if item.admin_level == 'village']
        province = next(item for item in record_repository.saved if item.admin_level == 'province')
        city = next(item for item in record_repository.saved if item.admin_level == 'city_regency')
        district = next(item for item in record_repository.saved if item.admin_level == 'district')

        assert province.external_code == '32'
        assert city.external_code == '3205'
        assert district.external_code == '3205160'
        assert len(villages) == 2
        assert all(item.external_code is None for item in villages)


class StubVersionRepositoryForWorkspace:
    def __init__(self, draft_version=None, published_version=None):
        self.draft_version = draft_version
        self.published_version = published_version

    def get_draft_version(self, registry_id):
        return self.draft_version if self.draft_version and self.draft_version.registry_id == registry_id else None

    def get_published_version(self, registry_id):
        return self.published_version if self.published_version and self.published_version.registry_id == registry_id else None


class StubVersionServiceForWorkspace:
    def __init__(self, draft_version=None, published_version=None):
        self.repository = StubVersionRepositoryForWorkspace(draft_version=draft_version, published_version=published_version)


class StubWorkspaceRecordRepository(SaveMixin):
    def __init__(self, record=None, records=None):
        super().__init__()
        self.record = record
        self.records = records or ([] if record is None else [record])

    def get_by_id(self, record_id):
        if self.record and self.record.id == record_id:
            return self.record
        for item in self.records:
            if item.id == record_id:
                return item
        return None

    def list_by_registry_version(self, registry_version_id, limit=None):
        items = [item for item in self.records if item.registry_version_id == registry_version_id]
        if limit is not None:
            items = items[:limit]
        return items

    def count_by_registry_version(self, registry_version_id):
        return len([item for item in self.records if item.registry_version_id == registry_version_id])


class StubBatchRepositoryForWorkspace:
    def list_by_registry_version(self, registry_version_id):
        return []


class StubImportBatchServiceForWorkspace:
    def extract_importable_fields(self, version):
        return version.schema_json.get('fields', [])


def test_get_registry_record_list_serializes_grid_columns_and_rows_for_generic_registry():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryService

        registry = DataRegistry(
            id=10,
            registry_slug='master.program',
            registry_code='MASTER-PROGRAM',
            name='Master Program',
            registry_type='master_data',
            category_key='master_data',
            source_mode='import_file',
            data_shape='hierarchical',
            status='draft',
        )
        draft_version = DataRegistryVersion(
            id=31,
            registry_id=10,
            version_number=2,
            schema_json={
                'fields': [
                    {'key': 'record_key', 'label': 'Record Key', 'type': 'text', 'required': True},
                    {'key': 'record_code', 'label': 'Kode Record', 'type': 'text', 'required': True},
                    {'key': 'label', 'label': 'Nama Program', 'type': 'text', 'required': True},
                    {'key': 'tipe', 'label': 'Tipe', 'type': 'select', 'required': False, 'options': ['D', 'E']},
                    {'key': 'jenis', 'label': 'Jenis', 'type': 'text', 'required': False},
                ]
            },
            status='draft',
            freshness_status='fresh',
        )
        records = [
            make_record(
                id=501,
                registry_id=10,
                registry_version_id=31,
                record_key='master_data:A1',
                record_code='A1',
                label='Program A1',
                admin_level='master_data',
                admin_level_code='MASTER_DATA',
                payload={'kode': 'A1', 'tipe': 'D', 'jenis': 'A'},
            ),
            make_record(
                id=502,
                registry_id=10,
                registry_version_id=31,
                record_key='master_data:B2',
                record_code='B2',
                label='Program B2',
                admin_level='master_data',
                admin_level_code='MASTER_DATA',
                is_active=False,
                payload={'kode': 'B2', 'tipe': 'E', 'jenis': 'B'},
            ),
        ]
        service = DataRegistryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_service=StubVersionServiceForWorkspace(draft_version=draft_version),
            record_repository=StubWorkspaceRecordRepository(records=records),
            batch_repository=StubBatchRepositoryForWorkspace(),
            import_batch_service=StubImportBatchServiceForWorkspace(),
        )

        result = service.get_registry_record_list(10, record_limit=100)

        assert result['record_count'] == 2
        assert [column['label'] for column in result['record_columns']] == ['Kode Record', 'Nama Program', 'Tipe', 'Jenis']
        assert result['edit_fields'][0]['key'] == 'record_code'
        assert result['edit_fields'][1]['key'] == 'label'
        assert result['record_rows'][0]['record_key'] == 'master_data:A1'
        assert result['record_rows'][0]['column_values']['tipe'] == 'D'
        assert result['record_rows'][1]['column_values']['jenis'] == 'B'
        assert result['record_rows'][1]['is_active'] is False


def test_update_registry_record_updates_generic_payload_and_identity_fields():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryService

        registry = DataRegistry(id=10, registry_slug='master.program', name='Master Program', status='draft')
        version = DataRegistryVersion(
            id=31,
            registry_id=10,
            version_number=2,
            schema_json={
                'fields': [
                    {'key': 'record_code', 'label': 'Kode Record', 'type': 'text', 'required': True},
                    {'key': 'label', 'label': 'Nama Program', 'type': 'text', 'required': True},
                    {'key': 'tipe', 'label': 'Tipe', 'type': 'select', 'required': False, 'options': ['D', 'E']},
                    {'key': 'jenis', 'label': 'Jenis', 'type': 'text', 'required': False},
                ]
            },
            status='draft',
            freshness_status='fresh',
        )
        record = make_record(
            id=501,
            registry_id=10,
            registry_version_id=31,
            record_key='master_data:A1',
            record_code='A1',
            label='Program A1',
            display_label='Program A1',
            normalized_label='program a1',
            admin_level='master_data',
            admin_level_code='MASTER_DATA',
            payload={'record_code': 'A1', 'label': 'Program A1', 'tipe': 'D', 'jenis': 'A'},
            registry_version=version,
        )
        actor = SimpleNamespace(id=7, uuid='actor-uuid')
        record_repository = StubWorkspaceRecordRepository(record=record)
        service = DataRegistryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_service=StubVersionServiceForWorkspace(draft_version=version),
            record_repository=record_repository,
            batch_repository=StubBatchRepositoryForWorkspace(),
            import_batch_service=StubImportBatchServiceForWorkspace(),
        )

        result = service.update_registry_record(
            501,
            {
                'registry_id': '10',
                'record_code': 'A1-REV',
                'label': 'Program A1 Revisi',
                'tipe': 'E',
                'jenis': 'A+',
            },
            actor=actor,
        )

        assert result['record'] is record
        assert record.record_code == 'A1-REV'
        assert record.label == 'Program A1 Revisi'
        assert record.display_label == 'Program A1 Revisi'
        assert record.normalized_label == 'program a1 revisi'
        assert record.payload['record_code'] == 'A1-REV'
        assert record.payload['label'] == 'Program A1 Revisi'
        assert record.payload['tipe'] == 'E'
        assert record.payload['jenis'] == 'A+'
        assert record.updated_by == 7
        assert record.updated_by_uuid == 'actor-uuid'
        assert record_repository.saved[-1] is record


def test_set_registry_record_active_updates_record_status():
    app = create_app('testing')

    with app.app_context():
        from app.modules.data_registry.services import DataRegistryService

        registry = DataRegistry(id=10, registry_slug='master.program', name='Master Program', status='draft')
        version = DataRegistryVersion(
            id=31,
            registry_id=10,
            version_number=2,
            schema_json={'fields': []},
            status='draft',
            freshness_status='fresh',
        )
        record = make_record(
            id=501,
            registry_id=10,
            registry_version_id=31,
            record_key='master_data:A1',
            record_code='A1',
            label='Program A1',
            admin_level='master_data',
            admin_level_code='MASTER_DATA',
            is_active=True,
            registry_version=version,
        )
        actor = SimpleNamespace(id=8, uuid='actor-uuid-2')
        record_repository = StubWorkspaceRecordRepository(record=record)
        service = DataRegistryService(
            registry_repository=StubDataRegistryRepository(registry=registry),
            version_service=StubVersionServiceForWorkspace(draft_version=version),
            record_repository=record_repository,
            batch_repository=StubBatchRepositoryForWorkspace(),
            import_batch_service=StubImportBatchServiceForWorkspace(),
        )

        result = service.set_registry_record_active(501, False, registry_id=10, actor=actor)

        assert result['record'] is record
        assert record.is_active is False
        assert record.updated_by == 8
        assert record.updated_by_uuid == 'actor-uuid-2'
        assert record_repository.saved[-1] is record

