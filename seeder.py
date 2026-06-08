import argparse
import csv
import os
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from dotenv import load_dotenv

from app.core.access import enrich_scope
from app.core.extensions import db
from app.modules.analytics.models import (
    AnalyticsIndicatorResult,
    AnalyticsReportVersionIndicator,
)
from app.modules.analytics.repositories import (
    AnalyticsDatasetRepository,
    AnalyticsDatasetVersionRepository,
    AnalyticsIndicatorDefinitionRepository,
    AnalyticsIndicatorResultRepository,
    AnalyticsIndicatorVersionRepository,
    AnalyticsReportDefinitionRepository,
    AnalyticsReportVersionRepository,
)
from app.modules.analytics.services import (
    AnalyticsDatasetRunService,
    AnalyticsDatasetService,
    AnalyticsIndicatorResultService,
    AnalyticsIndicatorService,
    AnalyticsReportService,
)
from app.modules.data_registry.repositories import DataRegistryRepository, DataRegistryVersionRepository
from app.modules.data_registry.services import (
    DataRegistryMaterializationService,
    DataRegistryService,
    DataRegistryVersionService,
)
from app.modules.employee.services import EmployeeService
from app.modules.form.models import Form
from app.modules.form.repositories import FormRepository
from app.modules.form.services import FormService, FormVersionService
from app.modules.submission.models import Submission, SubmissionEvent
from app.modules.submission.services import SubmissionService
from app.modules.user.repositories import UserRepository
from app.modules.user.services import UserService

load_dotenv()

DEFAULT_MATRIX_TEST_PASSWORD = 'MatrixTest#2026'
DEFAULT_WILAYAH_REGISTRY_SLUG = 'wilayah.administratif'
DEFAULT_WILAYAH_REGISTRY_CODE = 'WILAYAH-ADM'
DEFAULT_WILAYAH_REGISTRY_CSV_PATH = 'instance/diskominfo-od_kode_wilayah_dan_nama_wilayah_desa_kelurahan_data.csv'
DEFAULT_HARMONISASI_XLSX_PATH = 'instance/harmonisasi.xlsx'
DEFAULT_HARMONISASI_FORM_CODE = 'ANL-HARMONISASI-DEMO'
DEFAULT_HARMONISASI_FORM_SLUG = 'analytics-demo-harmonisasi-peraturan'
DEFAULT_HARMONISASI_DATASET_KEY = 'analytics-demo-harmonisasi-submission-fact'
DEFAULT_HARMONISASI_REPORT_KEY = 'analytics-demo-harmonisasi-overview'
DEFAULT_HARMONISASI_INDICATOR_KEY = 'analytics-demo-harmonisasi-selesai'
DEFAULT_HARMONISASI_SOURCE_TYPE = 'analytics_demo_harmonisasi'


def build_wilayah_registry_fixture_payload(csv_path=DEFAULT_WILAYAH_REGISTRY_CSV_PATH):
    return {
        'registry_slug': DEFAULT_WILAYAH_REGISTRY_SLUG,
        'registry_code': DEFAULT_WILAYAH_REGISTRY_CODE,
        'name': 'Wilayah Administratif',
        'description': 'Registry wilayah administratif hasil bootstrap dari CSV Diskominfo Jabar.',
        'registry_type': 'geo',
        'category_key': 'wilayah',
        'source_mode': 'import_file',
        'data_shape': 'hierarchical_geo',
        'schema_json': {
            'fields': [
                {'key': 'record_code'},
                {'key': 'label'},
                {'key': 'admin_level'},
                {'key': 'parent_record_code'},
            ]
        },
        'mapping_spec': {
            'source_format': 'csv',
            'delimiter': ',',
            'levels': ['province', 'city_regency', 'district', 'village'],
        },
        'source_snapshot': {
            'source_name': 'diskominfo-jabar',
            'source_file': csv_path,
        },
        'csv_path': csv_path,
    }


def seed_wilayah_registry(
    csv_path=DEFAULT_WILAYAH_REGISTRY_CSV_PATH,
    actor=None,
    registry_service=None,
    registry_repository=None,
    version_service=None,
    version_repository=None,
    materialization_service=None,
):
    fixture = build_wilayah_registry_fixture_payload(csv_path=csv_path)
    registry_service = registry_service or DataRegistryService()
    registry_repository = registry_repository or DataRegistryRepository()
    version_service = version_service or DataRegistryVersionService()
    version_repository = version_repository or DataRegistryVersionRepository()
    materialization_service = materialization_service or DataRegistryMaterializationService()

    if not os.path.exists(csv_path):
        raise FileNotFoundError(f'File CSV wilayah tidak ditemukan: {csv_path}')

    with open(csv_path, mode='r', encoding='utf-8', newline='') as file:
        source_rows = list(csv.DictReader(file))

    fixture['source_snapshot'] = {
        **fixture['source_snapshot'],
        'row_count': len(source_rows),
    }

    registry = registry_repository.get_by_slug(fixture['registry_slug'])
    if not registry:
        create_result = registry_service.create_registry(
            {
                'registry_slug': fixture['registry_slug'],
                'registry_code': fixture['registry_code'],
                'name': fixture['name'],
                'description': fixture['description'],
                'registry_type': fixture['registry_type'],
                'category_key': fixture['category_key'],
                'source_mode': fixture['source_mode'],
                'data_shape': fixture['data_shape'],
                'schema_json': fixture['schema_json'],
                'mapping_spec': fixture['mapping_spec'],
                'source_snapshot': fixture['source_snapshot'],
            },
            actor=actor,
        )
        registry = create_result['registry']
        draft_version = create_result['draft_version']
    else:
        registry.registry_code = fixture['registry_code']
        registry.name = fixture['name']
        registry.description = fixture['description']
        registry.registry_type = fixture['registry_type']
        registry.category_key = fixture['category_key']
        registry.source_mode = fixture['source_mode']
        registry.data_shape = fixture['data_shape']
        registry.schema_meta = {
            **(registry.schema_meta or {}),
            'source_name': fixture['source_snapshot']['source_name'],
            'source_file': fixture['source_snapshot']['source_file'],
        }
        registry_repository.save(registry)

        published_version = version_repository.get_published_version(registry.id)
        draft_version = version_repository.get_draft_version(registry.id)
        if not draft_version:
            draft_version = version_service.create_draft_version(
                registry_id=registry.id,
                schema_json=fixture['schema_json'],
                mapping_spec=fixture['mapping_spec'],
                source_snapshot=fixture['source_snapshot'],
                actor=actor,
                source_version_id=getattr(published_version, 'id', None),
            )

    draft_version.schema_json = fixture['schema_json']
    draft_version.mapping_spec = fixture['mapping_spec']
    draft_version.source_snapshot = fixture['source_snapshot']
    version_repository.save(draft_version)

    materialize_result = materialization_service.materialize_wilayah_rows(
        draft_version.id,
        source_rows,
        actor=actor,
    )
    publish_result = registry_service.publish_registry(
        registry.id,
        version_id=draft_version.id,
        actor=actor,
    )

    return {
        'registry': registry,
        'draft_version': draft_version,
        'published_version': publish_result['published_version'],
        'record_count': materialize_result['record_count'],
        'source_row_count': materialize_result['source_row_count'],
        'csv_path': csv_path,
    }


def print_wilayah_registry_summary(result):
    print('✅ Seed registry wilayah selesai.')
    print(
        f"- registry_slug={result['registry'].registry_slug} "
        f"published_version_id={getattr(result['published_version'], 'id', None)} "
        f"records={result['record_count']} rows={result['source_row_count']}"
    )
    print(f"- source_csv={result['csv_path']}")


def _normalize_harmonisasi_cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _normalize_harmonisasi_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()

    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return (datetime(1899, 12, 30) + timedelta(days=float(text))).date()
    except ValueError:
        return None


def load_harmonisasi_rows(xlsx_path=DEFAULT_HARMONISASI_XLSX_PATH):
    workbook = load_workbook(filename=xlsx_path, read_only=True, data_only=True)
    sheet = workbook['harmonisasi'] if 'harmonisasi' in workbook.sheetnames else workbook[workbook.sheetnames[0]]

    row_iterator = sheet.iter_rows(values_only=True)
    headers = [str(cell).strip() if cell is not None else '' for cell in next(row_iterator)]
    rows = []
    for excel_row_number, row_values in enumerate(row_iterator, start=2):
        if not any(value not in (None, '') for value in row_values):
            continue

        record = {
            header: _normalize_harmonisasi_cell(value)
            for header, value in zip(headers, row_values)
        }
        normalized_date = _normalize_harmonisasi_date(row_values[0] if row_values else None)
        record['TanggalISO'] = normalized_date.isoformat() if normalized_date else None
        record['TanggalDisplay'] = normalized_date.strftime('%d-%m-%Y') if normalized_date else ''
        record['__source_row_number'] = excel_row_number
        rows.append(record)

    workbook.close()
    return rows


def build_harmonisasi_demo_payload(xlsx_path=DEFAULT_HARMONISASI_XLSX_PATH):
    rows = load_harmonisasi_rows(xlsx_path=xlsx_path)
    owner_scope = enrich_scope({'type': 'kanwil', 'code': 'kanwil-jabar'})

    daerah_counter = Counter(row.get('Daerah') or 'Tanpa Daerah' for row in rows)
    jenis_counter = Counter(row.get('Jenis') or 'Tanpa Jenis' for row in rows)
    hasil_counter = Counter(row.get('Hasil') or 'Tanpa Hasil' for row in rows)
    metode_counter = Counter((row.get('Metode') or '').strip() or '(kosong)' for row in rows)
    tim_kerja_counter = Counter(row.get('Tim Kerja') or 'Tanpa Tim' for row in rows)

    total_rows = len(rows)
    selesai_count = hasil_counter.get('Selesai', 0)
    returned_count = hasil_counter.get('Dikembalikan', 0)
    achievement_percentage = round((selesai_count / total_rows) * 100, 2) if total_rows else 0
    reporting_year = max((
        _normalize_harmonisasi_date(row.get('TanggalISO')).year
        for row in rows
        if row.get('TanggalISO')
    ), default=datetime.now().year)

    row_snapshots = []
    yearly_periods = defaultdict(lambda: {'label': None, 'period_type': 'yearly', 'period_number': 1, 'reporting_year': None, 'total': 0, 'selesai': 0, 'dikembalikan': 0})
    semester_periods = defaultdict(lambda: {'label': None, 'period_type': 'semesterly', 'period_number': None, 'reporting_year': None, 'total': 0, 'selesai': 0, 'dikembalikan': 0})
    quarterly_periods = defaultdict(lambda: {'label': None, 'period_type': 'quarterly', 'period_number': None, 'reporting_year': None, 'total': 0, 'selesai': 0, 'dikembalikan': 0})

    for row in rows:
        iso_date = row.get('TanggalISO')
        normalized_date = _normalize_harmonisasi_date(iso_date) if iso_date else None
        row_year = normalized_date.year if normalized_date else reporting_year
        quarter_number = ((normalized_date.month - 1) // 3 + 1) if normalized_date else 1
        semester_number = 1 if (not normalized_date or normalized_date.month <= 6) else 2
        quarter_label = f'{row_year}-Q{quarter_number}'
        semester_label = f'{row_year}-S{semester_number}'
        year_label = str(row_year)
        snapshot = {
            'tanggal': iso_date,
            'reporting_year': row_year,
            'quarter_number': quarter_number,
            'quarter_label': quarter_label,
            'semester_number': semester_number,
            'semester_label': semester_label,
            'year_label': year_label,
            'daerah': row.get('Daerah') or 'Tanpa Daerah',
            'jenis': row.get('Jenis') or 'Tanpa Jenis',
            'hasil': row.get('Hasil') or 'Tanpa Hasil',
            'tim_kerja': row.get('Tim Kerja') or 'Tanpa Tim',
            'metode': (row.get('Metode') or '').strip() or '(kosong)',
        }
        row_snapshots.append(snapshot)

        for bucket, label, period_type, period_number in (
            (yearly_periods[year_label], year_label, 'yearly', 1),
            (semester_periods[semester_label], semester_label, 'semesterly', semester_number),
            (quarterly_periods[quarter_label], quarter_label, 'quarterly', quarter_number),
        ):
            bucket['label'] = label
            bucket['period_type'] = period_type
            bucket['period_number'] = period_number
            bucket['reporting_year'] = row_year
            bucket['total'] += 1
            if snapshot['hasil'] == 'Selesai':
                bucket['selesai'] += 1
            if snapshot['hasil'] == 'Dikembalikan':
                bucket['dikembalikan'] += 1

    def _serialize_period_metrics(periods):
        serialized = []
        for item in periods.values():
            total = item['total']
            serialized.append({
                **item,
                'achievement_percentage': round((item['selesai'] / total) * 100, 2) if total else 0,
            })
        return sorted(serialized, key=lambda item: (item['reporting_year'], item['period_number'], item['label']))

    summary = {
        'source_file': xlsx_path,
        'row_count': total_rows,
        'reporting_year': reporting_year,
        'date_min': min((row.get('TanggalISO') for row in rows if row.get('TanggalISO')), default=None),
        'date_max': max((row.get('TanggalISO') for row in rows if row.get('TanggalISO')), default=None),
        'jenis_counts': dict(jenis_counter),
        'hasil_counts': dict(hasil_counter),
        'metode_counts': dict(metode_counter),
        'tim_kerja_counts': dict(tim_kerja_counter),
        'top_daerah': [
            {'label': label, 'count': count}
            for label, count in daerah_counter.most_common(10)
        ],
        'completion': {
            'selesai': selesai_count,
            'dikembalikan': returned_count,
            'achievement_percentage': achievement_percentage,
        },
        'filter_dimensions': {
            'reporting_years': sorted({item['reporting_year'] for item in row_snapshots}),
            'hasil_options': sorted({item['hasil'] for item in row_snapshots}),
            'jenis_options': sorted({item['jenis'] for item in row_snapshots}),
        },
        'period_metrics': {
            'yearly': _serialize_period_metrics(yearly_periods),
            'semesterly': _serialize_period_metrics(semester_periods),
            'quarterly': _serialize_period_metrics(quarterly_periods),
        },
        'row_snapshots': row_snapshots,
    }

    submissions = []
    for row in rows:
        source_row_number = row['__source_row_number']
        submissions.append({
            'reporting_year': reporting_year,
            'payload': {
                'tanggal': row.get('TanggalISO') or row.get('TanggalDisplay'),
                'daerah': row.get('Daerah'),
                'jenis': row.get('Jenis'),
                'program_pembentukan': row.get('Program Pembentukan'),
                'urusan_pemerintahan': row.get('Urusan Pemerintahan'),
                'nama_raperda_raperkada': row.get('Nama Raperda/ Raperkada'),
                'tim_kerja': row.get('Tim Kerja'),
                'hasil': row.get('Hasil'),
                'metode': row.get('Metode') or None,
            },
            'context': {
                'meta': {
                    'source_file': xlsx_path,
                    'sheet_name': 'harmonisasi',
                    'source_row_number': source_row_number,
                },
                'source_type': DEFAULT_HARMONISASI_SOURCE_TYPE,
                'source_ref': f'harmonisasi:{source_row_number}',
                'owner_scope': owner_scope,
            },
        })

    first_top_daerah = summary['top_daerah'][0] if summary['top_daerah'] else {'label': '-', 'count': 0}
    first_preview_rows = [
        {
            'tanggal': row.get('TanggalISO'),
            'daerah': row.get('Daerah'),
            'jenis': row.get('Jenis'),
            'hasil': row.get('Hasil'),
            'tim_kerja': row.get('Tim Kerja'),
        }
        for row in rows[:5]
    ]

    return {
        'summary': summary,
        'submissions': submissions,
        'form': {
            'key': 'analytics_demo_harmonisasi',
            'code': DEFAULT_HARMONISASI_FORM_CODE,
            'slug': DEFAULT_HARMONISASI_FORM_SLUG,
            'name': 'Demo Harmonisasi Peraturan Perundang-undangan',
            'description': 'Form demo untuk mematerialisasi data harmonisasi dari workbook mini-lab ke submission store.',
            'owner_scope': owner_scope,
            'target_scope': owner_scope,
            'access_policy_key': 'form.kanwil_broadcast',
            'schema': {
                'display': 'form',
                'components': [
                    {'type': 'datetime', 'key': 'tanggal', 'label': 'Tanggal', 'input': True, 'validate': {'required': True}},
                    {'type': 'textfield', 'key': 'daerah', 'label': 'Daerah', 'input': True, 'validate': {'required': True}},
                    {'type': 'textfield', 'key': 'jenis', 'label': 'Jenis', 'input': True, 'validate': {'required': True}},
                    {'type': 'textfield', 'key': 'program_pembentukan', 'label': 'Program Pembentukan', 'input': True},
                    {'type': 'textarea', 'key': 'urusan_pemerintahan', 'label': 'Urusan Pemerintahan', 'input': True},
                    {'type': 'textarea', 'key': 'nama_raperda_raperkada', 'label': 'Nama Raperda/Raperkada', 'input': True, 'validate': {'required': True}},
                    {'type': 'textfield', 'key': 'tim_kerja', 'label': 'Tim Kerja', 'input': True},
                    {'type': 'select', 'key': 'hasil', 'label': 'Hasil', 'input': True, 'data': {'values': [{'label': 'Selesai', 'value': 'Selesai'}, {'label': 'Dikembalikan', 'value': 'Dikembalikan'}]}},
                    {'type': 'select', 'key': 'metode', 'label': 'Metode', 'input': True, 'data': {'values': [{'label': 'One Day Service', 'value': 'One Day Service'}, {'label': 'Normal', 'value': 'Normal'}, {'label': '(kosong)', 'value': ''}]}},
                    {'type': 'button', 'key': 'submit', 'label': 'Kirim', 'action': 'submit', 'input': True},
                ],
            },
        },
        'dataset': {
            'dataset_key': DEFAULT_HARMONISASI_DATASET_KEY,
            'name': 'Demo Dataset Harmonisasi Submission Fact',
            'description': 'Dataset demo Domain 5 yang merangkum submission harmonisasi peraturan dari workbook mini-lab.',
            'source_domain': 'submission',
            'source_type': 'aggregated_submission_fact',
            'primary_source_ref': DEFAULT_HARMONISASI_FORM_CODE,
            'status': 'active',
            'is_active': True,
            'is_year_scoped': True,
            'default_reporting_year_mode': 'explicit',
            'owner_scope_type': owner_scope.get('type'),
            'owner_scope_code': owner_scope.get('code'),
            'owner_scope_name': owner_scope.get('name'),
            'owner_scope_path': owner_scope.get('path'),
            'settings_json': {'form_code': DEFAULT_HARMONISASI_FORM_CODE, 'source_file': xlsx_path},
            'tags_json': ['demo', 'harmonisasi', 'analytics'],
        },
        'dataset_version': {
            'status': 'published',
            'is_current_draft': False,
            'is_current_published': True,
            'source_contract_json': {'form_code': DEFAULT_HARMONISASI_FORM_CODE, 'source_type': DEFAULT_HARMONISASI_SOURCE_TYPE, 'reporting_year': reporting_year},
            'query_spec_json': {'group_by': ['daerah', 'jenis', 'hasil', 'tim_kerja'], 'reporting_year': reporting_year},
            'transform_spec_json': {'derived_fields': ['tanggal', 'hasil', 'jenis', 'tim_kerja'], 'notes': 'Demo preview di-seed dari workbook harmonisasi.xlsx.'},
            'grain_key': 'per_submission',
            'output_schema_json': [{'key': 'tanggal', 'type': 'date'}, {'key': 'daerah', 'type': 'string'}, {'key': 'jenis', 'type': 'string'}, {'key': 'hasil', 'type': 'string'}, {'key': 'tim_kerja', 'type': 'string'}],
            'dimension_definitions_json': [{'key': 'daerah', 'label': 'Daerah'}, {'key': 'jenis', 'label': 'Jenis'}, {'key': 'hasil', 'label': 'Hasil'}, {'key': 'tim_kerja', 'label': 'Tim Kerja'}],
            'metric_definitions_json': [{'key': 'jumlah_dokumen', 'label': 'Jumlah Dokumen', 'aggregation': 'count'}],
            'default_filters_json': {'reporting_year': reporting_year},
            'sort_spec_json': [{'key': 'tanggal', 'direction': 'asc'}],
            'freshness_source_type': 'submissions.submitted_at',
            'freshness_source_ref': DEFAULT_HARMONISASI_FORM_CODE,
            'freshness_strategy': 'manual_assertion',
            'freshness_policy_json': {'expected_reporting_year': reporting_year},
            'publish_notes': 'Version seeded otomatis dari workbook harmonisasi mini-lab.',
        },
        'dataset_run': {
            'requested_reporting_year': reporting_year,
            'requested_filters_json': {'reporting_year': reporting_year, 'source_type': DEFAULT_HARMONISASI_SOURCE_TYPE},
            'freshness_status': 'fresh',
            'source_snapshot_json': {'source_file': xlsx_path, 'source_type': DEFAULT_HARMONISASI_SOURCE_TYPE, 'row_count': total_rows, 'date_range': {'start': summary['date_min'], 'end': summary['date_max']}},
            'result_row_count': total_rows,
            'result_schema_json': [{'key': 'daerah', 'type': 'string'}, {'key': 'jenis', 'type': 'string'}, {'key': 'hasil', 'type': 'string'}, {'key': 'tim_kerja', 'type': 'string'}, {'key': 'jumlah', 'type': 'integer'}],
            'result_preview_json': first_preview_rows,
            'materialization_ref': f'demo://harmonisasi/{reporting_year}',
            'summary_json': summary,
        },
        'report': {
            'report_key': DEFAULT_HARMONISASI_REPORT_KEY,
            'name': 'Demo Overview Harmonisasi Peraturan',
            'description': 'Report overview untuk membaca distribusi dan outcome harmonisasi peraturan pada data demo.',
            'report_type': 'custom',
            'category_key': 'harmonisasi_demo',
            'status': 'active',
            'is_active': True,
            'settings_json': {'source_file': xlsx_path},
            'tags_json': ['demo', 'harmonisasi'],
        },
        'report_version': {
            'status': 'published',
            'is_current_draft': False,
            'is_current_published': True,
            'title': 'Overview Harmonisasi 2026',
            'meta_description': f"Dataset demo memuat {total_rows} submission harmonisasi periode {summary['date_min']} s.d. {summary['date_max']}. Daerah terbanyak: {first_top_daerah['label']} ({first_top_daerah['count']} dokumen).",
            'structure_json': {'sections': [{'key': 'overview', 'title': 'Overview'}, {'key': 'hasil', 'title': 'Outcome Harmonisasi'}]},
            'narrative_guidance_json': {'focus': ['volume dokumen', 'status hasil', 'dominasi daerah']},
        },
        'indicator': {
            'indicator_key': DEFAULT_HARMONISASI_INDICATOR_KEY,
            'indicator_code': 'ANL-HARM-01',
            'name': 'Persentase Harmonisasi Selesai',
            'description': 'Mengukur proporsi dokumen harmonisasi yang berstatus selesai dibanding total submission demo.',
            'source_mode': 'dataset_driven',
            'calculation_type': 'percentage',
            'target_source_type': 'manual_central_target',
            'status': 'active',
            'is_active': True,
            'settings_json': {'source_file': xlsx_path},
            'tags_json': ['demo', 'harmonisasi'],
        },
        'indicator_version': {
            'status': 'published',
            'is_current_draft': False,
            'is_current_published': True,
            'period_mode': 'yearly',
            'aggregation_strategy': 'last_value',
            'unit_label': 'persen',
            'meta_description': f"{selesai_count} dari {total_rows} dokumen harmonisasi pada workbook demo sudah berstatus selesai.",
            'formula_json': {'expression': 'selesai_count / total_count * 100', 'inputs': {'selesai_count': selesai_count, 'total_count': total_rows}},
            'target_config_json': {'target_value': 90, 'unit': 'persen', 'comparison': 'gte'},
            'narrative_guidance_json': {'highlights': [f"Dokumen selesai {selesai_count} dari total {total_rows}.", f"Dokumen dikembalikan {returned_count}."]},
            'threshold_rules_json': [{'label': 'baik', 'min': 90}, {'label': 'perlu perhatian', 'max': 89.99}],
        },
        'indicator_result': {
            'reporting_year': reporting_year,
            'status': 'published',
            'completion_status': 'completed',
            'measured_value': achievement_percentage,
            'target_value': 90,
            'achievement_percentage': round((achievement_percentage / 90) * 100, 2) if 90 else 0,
            'qualitative_summary': f"{selesai_count} dokumen harmonisasi selesai dari total {total_rows}. Daerah terbanyak adalah {first_top_daerah['label']}.",
            'constraint_notes': f"Masih ada {returned_count} dokumen berstatus dikembalikan yang bisa dijadikan fokus tindak lanjut.",
            'narrative_context_json': {'hasil_counts': dict(hasil_counter), 'top_daerah': summary['top_daerah'][:5]},
        },
    }


def _upsert_harmonisasi_dataset(spec, actor=None):
    repository = AnalyticsDatasetRepository()
    service = AnalyticsDatasetService(dataset_repository=repository)
    dataset = repository.get_by_key(spec['dataset_key'])
    if not dataset:
        return service.create_dataset(spec, actor=actor)
    dataset.name = spec['name']
    dataset.description = spec['description']
    dataset.source_domain = spec['source_domain']
    dataset.source_type = spec['source_type']
    dataset.primary_source_ref = spec['primary_source_ref']
    dataset.status = spec['status']
    dataset.is_active = spec['is_active']
    dataset.is_year_scoped = spec['is_year_scoped']
    dataset.default_reporting_year_mode = spec['default_reporting_year_mode']
    dataset.owner_scope_type = spec['owner_scope_type']
    dataset.owner_scope_code = spec['owner_scope_code']
    dataset.owner_scope_name = spec['owner_scope_name']
    dataset.owner_scope_path = spec['owner_scope_path']
    dataset.settings_json = spec['settings_json']
    dataset.tags_json = spec['tags_json']
    return repository.save(dataset)


def _upsert_harmonisasi_dataset_version(dataset_id, spec, actor=None):
    repository = AnalyticsDatasetVersionRepository()
    service = AnalyticsDatasetService(dataset_version_repository=repository)
    version = repository.get_published_version(dataset_id) or repository.get_draft_version(dataset_id)
    if not version:
        version = service.create_dataset_version(dataset_id, spec, actor=actor)
    version.status = spec['status']
    version.is_current_draft = spec['is_current_draft']
    version.is_current_published = spec['is_current_published']
    version.source_contract_json = spec['source_contract_json']
    version.query_spec_json = spec['query_spec_json']
    version.transform_spec_json = spec['transform_spec_json']
    version.grain_key = spec['grain_key']
    version.output_schema_json = spec['output_schema_json']
    version.dimension_definitions_json = spec['dimension_definitions_json']
    version.metric_definitions_json = spec['metric_definitions_json']
    version.default_filters_json = spec['default_filters_json']
    version.sort_spec_json = spec['sort_spec_json']
    version.freshness_source_type = spec['freshness_source_type']
    version.freshness_source_ref = spec['freshness_source_ref']
    version.freshness_strategy = spec['freshness_strategy']
    version.freshness_policy_json = spec['freshness_policy_json']
    version.publish_notes = spec['publish_notes']
    repository.save(version)
    if version.status != 'published' or not version.is_current_published:
        version = service.publish_dataset_version(version.id, actor=actor)
    return version


def _upsert_harmonisasi_report(spec, version_spec, actor=None):
    repository = AnalyticsReportDefinitionRepository()
    version_repository = AnalyticsReportVersionRepository()
    service = AnalyticsReportService(report_definition_repository=repository, report_version_repository=version_repository)
    report = repository.get_by_key(spec['report_key'])
    if not report:
        report = service.create_report_definition(spec, actor=actor)
    else:
        report.name = spec['name']
        report.description = spec['description']
        report.report_type = spec['report_type']
        report.category_key = spec['category_key']
        report.status = spec['status']
        report.is_active = spec['is_active']
        report.settings_json = spec['settings_json']
        report.tags_json = spec['tags_json']
        repository.save(report)
    version = version_repository.get_published_version(report.id) or version_repository.get_draft_version(report.id)
    if not version:
        version = service.create_report_version(report.id, version_spec, actor=actor)
    version.status = version_spec['status']
    version.is_current_draft = version_spec['is_current_draft']
    version.is_current_published = version_spec['is_current_published']
    version.title = version_spec['title']
    version.meta_description = version_spec['meta_description']
    version.structure_json = version_spec['structure_json']
    version.narrative_guidance_json = version_spec['narrative_guidance_json']
    version_repository.save(version)
    if version.status != 'published' or not version.is_current_published:
        version = service.publish_report_version(version.id, actor=actor)
    return report, version


def _upsert_harmonisasi_indicator(spec, version_spec, dataset_id, dataset_version_id, actor=None):
    repository = AnalyticsIndicatorDefinitionRepository()
    version_repository = AnalyticsIndicatorVersionRepository()
    service = AnalyticsIndicatorService(indicator_definition_repository=repository, indicator_version_repository=version_repository)
    indicator = repository.get_by_key(spec['indicator_key'])
    if not indicator:
        indicator = service.create_indicator_definition(spec, actor=actor)
    else:
        indicator.indicator_code = spec['indicator_code']
        indicator.name = spec['name']
        indicator.description = spec['description']
        indicator.source_mode = spec['source_mode']
        indicator.calculation_type = spec['calculation_type']
        indicator.target_source_type = spec['target_source_type']
        indicator.status = spec['status']
        indicator.is_active = spec['is_active']
        indicator.settings_json = spec['settings_json']
        indicator.tags_json = spec['tags_json']
        repository.save(indicator)
    version = version_repository.get_published_version(indicator.id) or version_repository.get_draft_version(indicator.id)
    version_payload = {**version_spec, 'dataset_id': dataset_id, 'dataset_version_id': dataset_version_id}
    if not version:
        version = service.create_indicator_version(indicator.id, version_payload, actor=actor)
    version.dataset_id = dataset_id
    version.dataset_version_id = dataset_version_id
    version.status = version_payload['status']
    version.is_current_draft = version_payload['is_current_draft']
    version.is_current_published = version_payload['is_current_published']
    version.period_mode = version_payload['period_mode']
    version.aggregation_strategy = version_payload['aggregation_strategy']
    version.unit_label = version_payload['unit_label']
    version.meta_description = version_payload['meta_description']
    version.formula_json = version_payload['formula_json']
    version.target_config_json = version_payload['target_config_json']
    version.narrative_guidance_json = version_payload['narrative_guidance_json']
    version.threshold_rules_json = version_payload['threshold_rules_json']
    version_repository.save(version)
    if version.status != 'published' or not version.is_current_published:
        version = service.publish_indicator_version(version.id, actor=actor)
    return indicator, version


def _upsert_harmonisasi_indicator_result(spec, indicator_version_id, dataset_run_id, actor=None):
    repository = AnalyticsIndicatorResultRepository()
    service = AnalyticsIndicatorResultService(result_repository=repository)
    existing = repository.get_latest_for_period(indicator_version_id=indicator_version_id, reporting_year=spec['reporting_year'], reporting_period_id=None)
    payload = {**spec, 'indicator_version_id': indicator_version_id, 'dataset_run_id': dataset_run_id}
    if not existing:
        return service.record_result(payload, actor=actor)
    existing.dataset_run_id = dataset_run_id
    existing.status = payload['status']
    existing.completion_status = payload['completion_status']
    existing.measured_value = payload['measured_value']
    existing.target_value = payload['target_value']
    existing.achievement_percentage = payload['achievement_percentage']
    existing.qualitative_summary = payload['qualitative_summary']
    existing.constraint_notes = payload['constraint_notes']
    existing.narrative_context_json = payload['narrative_context_json']
    existing.source_snapshot_json = {'source_mode': 'dataset_driven', 'dataset_run_id': dataset_run_id, 'indicator_version_id': indicator_version_id}
    return repository.save(existing)


def _replace_harmonisasi_demo_submissions(form_id, submissions, actor=None):
    existing_submissions = Submission.query.filter(Submission.form_id == form_id, Submission.source_type == DEFAULT_HARMONISASI_SOURCE_TYPE).all()
    if existing_submissions:
        submission_ids = [submission.id for submission in existing_submissions]
        SubmissionEvent.query.filter(SubmissionEvent.submission_id.in_(submission_ids)).delete(synchronize_session=False)
        Submission.query.filter(Submission.id.in_(submission_ids)).delete(synchronize_session=False)
        db.session.commit()
    service = SubmissionService()
    created = []
    for submission in submissions:
        created.append(service.submit(form_id=form_id, payload=submission['payload'], actor=actor, context=submission['context'], reporting_year=submission['reporting_year']))
    return created


def seed_harmonisasi_analytics_demo(xlsx_path=DEFAULT_HARMONISASI_XLSX_PATH, actor=None):
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f'File workbook harmonisasi tidak ditemukan: {xlsx_path}')
    fixture = build_harmonisasi_demo_payload(xlsx_path=xlsx_path)
    form_record = _upsert_form_fixture(fixture['form'], actor=actor, form_service=FormService(), form_repository=FormRepository(), form_version_service=FormVersionService())
    form = Form.query.filter_by(code=fixture['form']['code']).first()
    created_submissions = _replace_harmonisasi_demo_submissions(form.id, fixture['submissions'], actor=actor)
    dataset = _upsert_harmonisasi_dataset(fixture['dataset'], actor=actor)
    dataset_version = _upsert_harmonisasi_dataset_version(dataset.id, fixture['dataset_version'], actor=actor)
    run_service = AnalyticsDatasetRunService()
    run = run_service.start_run(dataset_id=dataset.id, dataset_version_id=dataset_version.id, data={'trigger_type': 'manual', 'trigger_ref': 'seeder:harmonisasi_demo', 'requested_reporting_year': fixture['dataset_run']['requested_reporting_year'], 'requested_filters_json': fixture['dataset_run']['requested_filters_json'], 'freshness_status': fixture['dataset_run']['freshness_status'], 'source_snapshot_json': fixture['dataset_run']['source_snapshot_json']}, actor=actor)
    run = run_service.complete_run(run.id, data={'result_row_count': fixture['dataset_run']['result_row_count'], 'result_schema_json': fixture['dataset_run']['result_schema_json'], 'result_preview_json': fixture['dataset_run']['result_preview_json'], 'materialization_ref': fixture['dataset_run']['materialization_ref'], 'summary_json': fixture['dataset_run']['summary_json']}, actor=actor)
    run.freshness_status = fixture['dataset_run']['freshness_status']
    run.source_snapshot_json = fixture['dataset_run']['source_snapshot_json']
    run_service.repository.save(run)
    report, report_version = _upsert_harmonisasi_report(fixture['report'], fixture['report_version'], actor=actor)
    indicator, indicator_version = _upsert_harmonisasi_indicator(fixture['indicator'], fixture['indicator_version'], dataset.id, dataset_version.id, actor=actor)
    mapping = AnalyticsReportVersionIndicator.query.filter_by(report_version_id=report_version.id, indicator_version_id=indicator_version.id).first()
    if not mapping:
        mapping = AnalyticsIndicatorService().attach_indicator_to_report_version(report_version_id=report_version.id, indicator_version_id=indicator_version.id, data={'item_order': 1, 'display_label': indicator.name, 'section_key': 'hasil', 'config_json': {'source': 'seed_demo'}}, actor=actor)
    result = _upsert_harmonisasi_indicator_result(fixture['indicator_result'], indicator_version_id=indicator_version.id, dataset_run_id=run.id, actor=actor)
    return {'source_file': xlsx_path, 'form': form_record, 'submission_count': len(created_submissions), 'dataset_id': dataset.id, 'dataset_version_id': dataset_version.id, 'dataset_run_id': run.id, 'report_id': report.id, 'report_version_id': report_version.id, 'indicator_id': indicator.id, 'indicator_version_id': indicator_version.id, 'indicator_result_id': result.id, 'indicator_mapping_id': mapping.id, 'summary': fixture['summary']}


def print_harmonisasi_analytics_demo_summary(result):
    completion = result['summary']['completion']
    print('✅ Seed demo analytics harmonisasi selesai.')
    print(f"- source_file={result['source_file']}")
    print(f"- form_code={result['form']['code']} submissions={result['submission_count']} dataset_id={result['dataset_id']} dataset_version_id={result['dataset_version_id']} run_id={result['dataset_run_id']}")
    print(f"- report_id={result['report_id']} indicator_id={result['indicator_id']} indicator_result_id={result['indicator_result_id']}")
    print(f"- periode={result['summary']['date_min']}..{result['summary']['date_max']} selesai={completion['selesai']} dikembalikan={completion['dikembalikan']} persentase_selesai={completion['achievement_percentage']}%")


def seed_employees():
    """Seed data pegawai dari CSV."""
    file_path = 'instance/pegawai.csv'
    emp_service = EmployeeService()

    with open(file_path, mode='r', encoding='utf-8') as file:
        reader = csv.DictReader(file, delimiter=';')
        for row in reader:
            try:
                emp_service.add(
                    nip=row['nip'],
                    name=row['full_name'],
                    unit=row['unit'],
                    skip_if_exists=True,
                )
                print(f"✅ Pegawai OK: {row['nip']}")
            except ValueError as error:
                print(f"❌ Error seed pegawai {row.get('nip')}: {str(error)}")
    print('✅ Seeding data pegawai selesai.')


def seed_user_agents():
    """Seed data user Agentic AI."""
    user = _upsert_user(
        {
            'key': 'agentic_ai',
            'username': '999999999999999999',
            'email': 'agent.super@kemenkum.go.id',
            'password': 'JalanJakarta#27',
            'roles': ['superadmin'],
            'permissions': [],
            'settings': {'active_year': 2026},
            'active': True,
        }
    )
    print(f'✅ Seeding data agen selesai: {user.username}')
    return user


def build_matrix_fixture_payload(password=DEFAULT_MATRIX_TEST_PASSWORD):
    admin_scope = enrich_scope({'type': 'kanwil', 'code': 'kanwil-jabar'})
    unit_ki_scope = enrich_scope({'type': 'unit', 'code': 'ki'})
    unit_ahu_scope = enrich_scope({'type': 'unit', 'code': 'ahu'})
    division_scope = enrich_scope({'type': 'division', 'code': 'divisi-pelayanan-hukum'})

    users = [
        {
            'key': 'admin_kanwil',
            'label': 'Admin Kanwil',
            'username': '888888888888888888',
            'email': 'agen.admin@kemenkum.go.id',
            'password': password,
            'roles': ['admin_lintas_bagian'],
            'permissions': [],
            'settings': {
                'active_year': 2026,
                'scope': admin_scope,
            },
            'active': True,
        },
        {
            'key': 'pegawai_ki',
            'label': 'Pegawai Unit KI',
            'username': '555555555555555555',
            'email': 'agen.ki@kemenkum.go.id',
            'password': password,
            'roles': ['pegawai_unit'],
            'permissions': ['submission:create'],
            'settings': {
                'active_year': 2026,
                'scope': unit_ki_scope,
            },
            'active': True,
        },
        {
            'key': 'pegawai_ahu',
            'label': 'Pegawai Unit AHU',
            'username': '666666666666666666',
            'email': 'agen.ahu@kemenkum.go.id',
            'password': password,
            'roles': ['pegawai_unit'],
            'permissions': ['submission:create'],
            'settings': {
                'active_year': 2026,
                'scope': unit_ahu_scope,
            },
            'active': True,
        },
        {
            'key': 'kadiv_yankum',
            'label': 'Kepala Divisi Pelayanan Hukum',
            'username': '777777777777777777',
            'email': 'agen.kadiv.yankum@kemenkum.go.id',
            'password': password,
            'roles': ['kepala_divisi'],
            'permissions': ['submission:create'],
            'settings': {
                'active_year': 2026,
                'scope': division_scope,
            },
            'active': True,
        },
    ]

    forms = [
        {
            'key': 'f1_ki',
            'code': 'MTRX-F1-KI',
            'slug': 'matrix-target-unit-ki',
            'name': 'Matrix F1 Target Unit KI',
            'description': 'Fixture manual test matriks ABAC untuk target unit KI.',
            'owner_scope': admin_scope,
            'target_scope': unit_ki_scope,
            'access_policy_key': 'form.unit_internal',
            'schema': {
                'display': 'form',
                'components': [
                    {
                        'type': 'textfield',
                        'key': 'judul_kegiatan',
                        'label': 'Judul Kegiatan',
                        'validate': {'required': True},
                        'input': True,
                    },
                    {
                        'type': 'textarea',
                        'key': 'catatan',
                        'label': 'Catatan',
                        'input': True,
                    },
                    {
                        'type': 'button',
                        'key': 'submit',
                        'label': 'Kirim',
                        'action': 'submit',
                        'input': True,
                    },
                ],
            },
        },
        {
            'key': 'f2_ahu',
            'code': 'MTRX-F2-AHU',
            'slug': 'matrix-target-unit-ahu',
            'name': 'Matrix F2 Target Unit AHU',
            'description': 'Fixture manual test matriks ABAC untuk target unit AHU.',
            'owner_scope': admin_scope,
            'target_scope': unit_ahu_scope,
            'access_policy_key': 'form.unit_internal',
            'schema': {
                'display': 'form',
                'components': [
                    {
                        'type': 'textfield',
                        'key': 'judul_kegiatan',
                        'label': 'Judul Kegiatan',
                        'validate': {'required': True},
                        'input': True,
                    },
                    {
                        'type': 'number',
                        'key': 'jumlah_layanan',
                        'label': 'Jumlah Layanan',
                        'input': True,
                    },
                    {
                        'type': 'button',
                        'key': 'submit',
                        'label': 'Kirim',
                        'action': 'submit',
                        'input': True,
                    },
                ],
            },
        },
        {
            'key': 'f3_divph',
            'code': 'MTRX-F3-DIVPH',
            'slug': 'matrix-target-divisi-pelayanan-hukum',
            'name': 'Matrix F3 Broadcast Divisi Pelayanan Hukum',
            'description': 'Fixture manual test matriks ABAC untuk target divisi Pelayanan Hukum.',
            'owner_scope': admin_scope,
            'target_scope': division_scope,
            'access_policy_key': 'form.division_broadcast',
            'schema': {
                'display': 'form',
                'components': [
                    {
                        'type': 'textfield',
                        'key': 'judul_kegiatan',
                        'label': 'Judul Kegiatan',
                        'validate': {'required': True},
                        'input': True,
                    },
                    {
                        'type': 'select',
                        'key': 'status_realisasi',
                        'label': 'Status Realisasi',
                        'input': True,
                        'data': {
                            'values': [
                                {'label': 'Belum Jalan', 'value': 'belum_jalan'},
                                {'label': 'Proses', 'value': 'proses'},
                                {'label': 'Selesai', 'value': 'selesai'},
                            ]
                        },
                    },
                    {
                        'type': 'button',
                        'key': 'submit',
                        'label': 'Kirim',
                        'action': 'submit',
                        'input': True,
                    },
                ],
            },
        },
    ]

    return {
        'password': password,
        'users': users,
        'forms': forms,
    }


def seed_matrix_test_fixture(
    password=DEFAULT_MATRIX_TEST_PASSWORD,
    user_service=None,
    user_repository=None,
    form_service=None,
    form_repository=None,
    form_version_service=None,
):
    fixture = build_matrix_fixture_payload(password=password)
    user_service = user_service or UserService()
    user_repository = user_repository or UserRepository()
    form_repository = form_repository or FormRepository()
    form_service = form_service or FormService(form_repository=form_repository)
    form_version_service = form_version_service or FormVersionService(form_repository=form_repository)

    seeded_users = []
    actor_map = {}
    for user_spec in fixture['users']:
        user = _upsert_user(user_spec, user_service=user_service, user_repository=user_repository)
        seeded_users.append(
            {
                'key': user_spec['key'],
                'label': user_spec['label'],
                'username': user.username,
                'email': user.email,
                'roles': list(user.roles or []),
                'permissions': list(user.permissions or []),
                'scope_code': ((user.settings or {}).get('scope') or {}).get('code'),
            }
        )
        actor_map[user_spec['key']] = user

    admin_actor = actor_map['admin_kanwil']
    seeded_forms = []
    for form_spec in fixture['forms']:
        form_record = _upsert_form_fixture(
            form_spec,
            actor=admin_actor,
            form_service=form_service,
            form_repository=form_repository,
            form_version_service=form_version_service,
        )
        seeded_forms.append(form_record)

    return {
        'password': password,
        'users': seeded_users,
        'forms': seeded_forms,
    }


def _upsert_user(spec, user_service=None, user_repository=None):
    user_service = user_service or UserService()
    user_repository = user_repository or UserRepository()
    existing_user = user_repository.find_one_by(username=spec['username'])

    if not existing_user:
        return user_service.register(
            username=spec['username'],
            email=spec['email'],
            password=spec['password'],
            is_seeding=True,
            roles=spec.get('roles'),
            permissions=spec.get('permissions'),
            active=spec.get('active', True),
            settings=spec.get('settings'),
        )

    existing_user.email = spec['email']
    existing_user.active = bool(spec.get('active', True))
    existing_user.roles = list(spec.get('roles') or [])
    existing_user.permissions = list(spec.get('permissions') or [])
    existing_user.settings = spec.get('settings') or {}
    existing_user.set_password(spec['password'])
    return user_repository.save(existing_user)


def _upsert_form_fixture(spec, actor, form_service=None, form_repository=None, form_version_service=None):
    form_repository = form_repository or FormRepository()
    form_service = form_service or FormService(form_repository=form_repository)
    form_version_service = form_version_service or FormVersionService(form_repository=form_repository)
    existing_form = form_repository.find_one_by(code=spec['code'])

    form_payload = {
        'code': spec['code'],
        'slug': spec['slug'],
        'name': spec['name'],
        'description': spec.get('description'),
        'visibility': 'internal',
        'owner_scope': spec.get('owner_scope'),
        'target_scope': spec.get('target_scope'),
        'access_policy_key': spec.get('access_policy_key'),
        'schema': spec.get('schema'),
    }

    if existing_form:
        form = form_service.update_form_identity(existing_form.id, form_payload, actor=actor)
    else:
        create_result = form_service.create_form(form_payload, actor=actor)
        form = create_result['form']

    draft_version = form_version_service.create_draft_version(
        form_id=form.id,
        schema=spec.get('schema'),
        actor=actor,
    )
    published_version = form_version_service.publish_version(draft_version.id, actor=actor)

    target_scope = spec.get('target_scope') or {}
    return {
        'key': spec['key'],
        'form_id': form.id,
        'code': spec['code'],
        'slug': spec['slug'],
        'name': spec['name'],
        'target_scope_type': target_scope.get('type'),
        'target_scope_code': target_scope.get('code'),
        'access_policy_key': spec.get('access_policy_key'),
        'published_version_id': getattr(published_version, 'id', None),
    }


def print_matrix_fixture_summary(result):
    print('✅ Seed fixture matriks ABAC selesai.')
    print(f"Password semua user fixture: {result['password']}")
    print('Users:')
    for user in result['users']:
        role_list = ', '.join(user['roles']) if user['roles'] else '-'
        permission_list = ', '.join(user['permissions']) if user['permissions'] else '-'
        print(
            f"- {user['label']}: username={user['username']} email={user['email']} "
            f"scope={user['scope_code']} roles={role_list} permissions={permission_list}"
        )
    print('Forms:')
    for form in result['forms']:
        print(
            f"- {form['code']} ({form['name']}): form_id={form['form_id']} "
            f"target={form['target_scope_type']}:{form['target_scope_code']} "
            f"policy={form['access_policy_key']} published_version_id={form['published_version_id']}"
        )


def parse_args():
    parser = argparse.ArgumentParser(description='Seeder proyek dasborkanwil')
    parser.add_argument(
        '--profile',
        action='append',
        choices=['employees', 'agent', 'abac_matrix', 'wilayah_registry', 'analytics_harmonisasi_demo', 'all'],
        help='Profile seeding yang akan dijalankan. Bisa dipakai berulang.',
    )
    parser.add_argument(
        '--matrix-password',
        default=DEFAULT_MATRIX_TEST_PASSWORD,
        help='Password yang dipakai semua user fixture matriks ABAC.',
    )
    parser.add_argument(
        '--wilayah-csv',
        default=DEFAULT_WILAYAH_REGISTRY_CSV_PATH,
        help='Path CSV bootstrap registry wilayah administratif.',
    )
    parser.add_argument(
        '--harmonisasi-xlsx',
        default=DEFAULT_HARMONISASI_XLSX_PATH,
        help='Path workbook demo harmonisasi untuk seed analytics.',
    )
    return parser.parse_args()


def run_selected_seed_profiles(args):
    selected_profiles = args.profile or ['all']
    if 'all' in selected_profiles:
        selected_profiles = ['employees', 'agent', 'abac_matrix', 'wilayah_registry', 'analytics_harmonisasi_demo']

    if 'employees' in selected_profiles:
        seed_employees()
    if 'agent' in selected_profiles:
        seed_user_agents()
    if 'abac_matrix' in selected_profiles:
        result = seed_matrix_test_fixture(password=args.matrix_password)
        print_matrix_fixture_summary(result)
    if 'wilayah_registry' in selected_profiles:
        result = seed_wilayah_registry(csv_path=args.wilayah_csv)
        print_wilayah_registry_summary(result)
    if 'analytics_harmonisasi_demo' in selected_profiles:
        result = seed_harmonisasi_analytics_demo(xlsx_path=args.harmonisasi_xlsx)
        print_harmonisasi_analytics_demo_summary(result)


if __name__ == '__main__':
    from app import create_app

    arguments = parse_args()
    app = create_app(config_mode=os.getenv('FLASK_ENV', 'default'))
    with app.app_context():
        run_selected_seed_profiles(arguments)
